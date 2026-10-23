# ======== agent v28j (18 Oct 2026): v28i plus a crash fix in search_titles (Wikipedia reply without 'query'), a safety net that turns unexpected errors into a short try-again message, and a mild note instead of a warning when the source text has no 'first'. Syntax-checked; parts tested live as patches. ========

# ======== agent v28i (18 Oct 2026): v28h plus a wider code-request rule (verbs compute/calculate/plot/print..., 'Using X, <verb>' openings, pandas/numpy/matplotlib/sqlite/unittest/flask, .py/.js names, calculations with numbers). Syntax-checked; tested live as a patch on 10 tasks (10/10), not yet as a standalone file. ========

# ======== agent v28h (18 Oct 2026): v28g plus the soft-mode evidence check (warning line, answer not saved when the source text has no 'first' or has an extra qualifier). Built on Kaggle; syntax-checked, not yet run as a standalone file. ========



# ======== # Shatachandra AI Lab agent v28g (Colab/Kaggle), built 17 Oct 2026 from v27d plus the experimental patches. Not yet tested as a standalone file. ========



# ======== v27d agent ========
# Kaggle agent v3, 12 Oct 2026. Not tested as a standalone file.
import re, math, time, json, io, contextlib, unicodedata, requests, torch
from bs4 import BeautifulSoup
H = {'User-Agent': 'shatachandra-ai-lab/0.1 (personal research)'}
API = 'https://en.wikipedia.org/w/api.php'
NOT_FOUND = {'Telugu': 'నేను దీనిని కనుగొనలేకపోయాను.', 'Hindi': 'मैं यह जानकारी नहीं ढूंढ सका।', 'English': 'I could not find it.'}
LOOKUP_PROMPT = 'You can use one tool: lookup(page), which opens an English Wikipedia page. If you need facts you are not sure about, reply ONLY with a line like:\nTOOL: lookup("Page Title")\nThe page title must be the exact English Wikipedia title of the person, place or thing. Do not answer from memory when you use the tool.\n\n'
SYN = {'built': ['built', 'constructed', 'commissioned', 'founded', 'established'], 'build': ['built', 'constructed', 'commissioned', 'founded', 'established'], 'flow': ['flow', 'flows', 'mouth', 'empties', 'drains', 'discharges'], 'wrote': ['wrote', 'written', 'author', 'authored', 'composed', 'lyrics'], 'capital': ['capital', 'seat']}

def nfc(s):
    return unicodedata.normalize("NFC", s).lower()


def is_hit(ans, expected):
    a = nfc(ans)
    for grp in expected:
        ok = False
        for k in grp:
            k = nfc(k)
            if k.isdigit():
                if re.search(r"(?<!\d)" + k + r"(?!\d)", a):
                    ok = True
            elif k in a:
                ok = True
        if not ok:
            return False
    return True


def is_safe(ans):
    a = ans.strip()
    return a in NOT_FOUND.values() or any(w in a.lower() for w in SAFE_WORDS)


def detect_lang(q):
    if re.search(r"[\u0C00-\u0C7F]", q): return "Telugu"
    if re.search(r"[\u0900-\u097F]", q): return "Hindi"
    return "English"


def ask_long(q, seed, n=400):
    torch.manual_seed(seed)
    msgs = [{"role": "user", "content": [{"type": "text", "text": q}]}]
    inputs = tokenizer.apply_chat_template(
        msgs, add_generation_prompt=True, tokenize=True,
        return_dict=True, return_tensors="pt",
    ).to("cuda")
    out = model.generate(**inputs, max_new_tokens=n, do_sample=True,
                         temperature=0.7, top_p=0.8, top_k=20)
    return tokenizer.decode(out[0][inputs["input_ids"].shape[1]:], skip_special_tokens=True)


def toks(s):
    return set(re.findall(r"\w+", s.lower()))


def wiki_get(params, tries=5):
    for i in range(tries):
        r = requests.get(API, params=params, headers=H, timeout=15)
        if r.status_code == 200 and "json" in r.headers.get("content-type", ""):
            return r.json()
        try:
            wait = float(r.headers.get("Retry-After", 2 * (i + 1)))
        except ValueError:
            wait = 2 * (i + 1)
        time.sleep(min(wait, 20))
    raise RuntimeError("Wikipedia kept refusing, last status " + str(r.status_code))


def page_blocks2(title):
    j = wiki_get({"action": "parse", "format": "json", "formatversion": 2,
                  "prop": "text", "page": title, "redirects": 1})
    soup = BeautifulSoup(j["parse"]["text"], "html.parser")
    for t in soup(["style", "script", "sup"]):
        t.decompose()
    blocks = []
    for el in soup.find_all(["p", "li", "tr"]):
        s = re.sub(r"\s+", " ", el.get_text(" ", strip=True))
        if len(s) > 15:
            blocks.append(s)
    return blocks


def exists(title):
    try:
        j = wiki_get({"action": "parse", "format": "json", "formatversion": 2,
                      "prop": "text", "page": title, "redirects": 1})
        return "parse" in j
    except Exception:
        return False


def search_titles(query, n):
    j = wiki_get({"action": "query", "format": "json", "list": "search",
                  "srsearch": query, "srlimit": n})
    return [x["title"] for x in j["query"]["search"]]


def pick_from_search(eng, seed, exclude=()):
    titles = [t for t in search_titles(eng, 8) if t not in exclude][:6]
    if not titles:
        return None
    menu = "\n".join(f"{i+1}. {t}" for i, t in enumerate(titles))
    pick = ask_long("Question: " + eng + "\n\nWikipedia pages:\n" + menu +
                    "\n\nReply with only the number of the page most likely to contain the answer.",
                    seed, n=5).strip()
    m = re.search(r"\d", pick)
    i = int(m.group()) - 1 if m else 0
    return titles[i if 0 <= i < len(titles) else 0]


def expand(query):
    words = set(w for w in toks(query) if len(w) > 2)
    extra = set()
    for w in words:
        for k, vs in SYN.items():
            if w.startswith(k) or w in vs:
                extra |= set(vs)
    return words | extra


def lookup(title, query, cap=6000):
    try:
        blocks = page_blocks2(title)
    except Exception:
        blocks = []
    blocks = [b for b in blocks if not b.startswith("^")]
    qw = expand(query) - {"the", "what", "who", "when", "where", "which", "how", "india"}
    lead = list(range(min(4, len(blocks))))
    short = [i for i in range(4, len(blocks)) if len(blocks[i]) < 60 and qw & toks(blocks[i])][:8]
    scored = sorted(range(4, len(blocks)), key=lambda i: -len(qw & toks(blocks[i])))
    chosen, total = [], 0
    for i in lead + short + scored:
        if i in chosen:
            continue
        s = blocks[i][:250]
        if total + len(s) > cap:
            break
        chosen.append(i)
        total += len(s)
    chosen.sort()
    return [blocks[i][:250] for i in chosen]


def agent(question, seed=0, max_tries=3):
    lang = detect_lang(question)
    answer_prompt = (
        "Answer the question ONLY from the RESULT below. "
        f"Write the answer in {lang}, in one or two short sentences. "
        "Keep the exact meaning of the RESULT. Do not say the RESULT lacks the answer if it contains it.\n\n"
    )
    tried = []
    with model.disable_adapter():
        eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                       seed, n=40).strip()
        print("ENGLISH:", eng)
        for attempt in range(max_tries):
            if attempt == 0:
                first = ask_long(LOOKUP_PROMPT + "Question: " + eng, seed, n=60).strip()
                print("TRY 1:", first)
                m = re.search(r'lookup\("(.+?)"\)', first)
                if not m:
                    return first
                title = m.group(1)
                if not exists(title):
                    title = pick_from_search(eng, seed, exclude=tried)
            else:
                title = pick_from_search(eng, seed, exclude=tried)
            if not title or title in tried:
                break
            tried.append(title)
            result = "\n".join(lookup(title, eng))
            check = ask_long("Question: " + eng + "\nRESULT: " + result +
                "\n\nDoes the RESULT contain the answer to the question? Reply with only YES or NO.",
                seed, n=5).strip().upper()
            print("  PAGE:", title, "| CHECK:", check)
            if check.startswith("YES"):
                return ask_long(answer_prompt + "Question: " + question + "\nRESULT: " + result,
                                seed, n=120).strip()
    return NOT_FOUND[lang]


# ---- v4 (13 Oct 2026): chunk scan replaces line picking. Not tested as a standalone file. ----
import time

def chunk_scan(title, eng, max_chars=4500, max_chunks=12):
    blocks = [b for b in page_blocks2(title) if not b.startswith("^")]
    chunks, cur = [], ""
    for b in blocks:
        if len(cur) + len(b) > max_chars:
            chunks.append(cur)
            cur = ""
        cur += b[:300] + "\n"
    if cur:
        chunks.append(cur)
    for i, c in enumerate(chunks[:max_chunks]):
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c +
                     "\n\nDoes this TEXT directly state the answer? If yes, reply 'YES: ' followed by the answer. "
                     "If not, reply 'NO'.", 0, n=40).strip()
        if r.upper().startswith("YES"):
            return r
    return None

def agent(question, seed=0, max_tries=3):
    lang = detect_lang(question)
    tried = []
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    first = ask_long(LOOKUP_PROMPT + "Question: " + eng, seed, n=60).strip()
    m = re.search(r'lookup\("(.+?)"\)', first)
    if not m:
        return first
    title = m.group(1)
    if not exists(title):
        title = pick_from_search(eng, seed, exclude=tried)
    for attempt in range(max_tries):
        if not title or title in tried:
            break
        tried.append(title)
        found = chunk_scan(title, eng)
        if found:
            facts = found[4:].strip() if found.upper().startswith("YES:") else found
            return ask_long(f"Question: {question}\nANSWER FROM A SOURCE: {facts}\n\n"
                            f"Write this answer in {lang}, in one or two short sentences. "
                            "Keep the exact meaning.", seed, n=120).strip()
        title = pick_from_search(eng, seed, exclude=tried)
    return NOT_FOUND[lang]


# ---- v5 (13 Oct 2026): second search with the model's own title, menu of 6, no translation hint, no-tool path returns direct answers. Not tested standalone. ----
def pick_from_search(eng, seed, exclude=(), extra_query=None):
    titles = []
    for qq in [eng] + ([extra_query] if extra_query else []):
        for t in search_titles(qq, 8):
            if t not in exclude and t not in titles:
                titles.append(t)
    titles = titles[:6]
    if not titles:
        return None
    menu = "\n".join(f"{i+1}. {t}" for i, t in enumerate(titles))
    pick = ask_long("Question: " + eng + "\n\nWikipedia pages:\n" + menu +
                    "\n\nReply with only the number of the page most likely to contain the answer.",
                    seed, n=5).strip()
    m = re.search(r"\d", pick)
    i = int(m.group()) - 1 if m else 0
    return titles[i if 0 <= i < len(titles) else 0]


def agent(question, seed=0, max_tries=3):
    lang = detect_lang(question)
    tried = []
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    first = ask_long(LOOKUP_PROMPT + "Question: " + eng, seed, n=60).strip()
    m = re.search(r'lookup\("(.+?)"\)', first)
    if not m:
        if lang == "English" or len(first) < 200 and not first.lower().startswith(("i need", "could you")):
            return first
        return NOT_FOUND[lang]
    title = m.group(1)
    own = title
    if not exists(title):
        title = pick_from_search(eng, seed, exclude=tried, extra_query=own)
    for attempt in range(max_tries):
        if not title or title in tried:
            break
        tried.append(title)
        found = chunk_scan(title, eng)
        if found:
            facts = found[4:].strip() if found.upper().startswith("YES:") else found
            return ask_long(f"Question: {question}\nANSWER FROM A SOURCE: {facts}\n\n"
                            f"Write this answer in {lang}, in one or two short sentences. "
                            "Keep the exact meaning.", seed, n=120).strip()
        title = pick_from_search(eng, seed, exclude=tried, extra_query=own)
    return NOT_FOUND[lang]



# ---- v6 (13 Oct 2026): menu of 6 first, 8 on retry; write_answer helper. Not tested standalone. ----
def pick_from_search(eng, seed, exclude=(), extra_query=None, size=6):
    titles = []
    for qq in [eng] + ([extra_query] if extra_query else []):
        for t in search_titles(qq, 8):
            if t not in exclude and t not in titles:
                titles.append(t)
    titles = titles[:size]
    if not titles:
        return None
    menu = "\n".join(f"{i+1}. {t}" for i, t in enumerate(titles))
    pick = ask_long("Question: " + eng + "\n\nWikipedia pages:\n" + menu +
                    "\n\nReply with only the number of the page most likely to contain the answer.",
                    seed, n=5).strip()
    m = re.search(r"\d", pick)
    i = int(m.group()) - 1 if m else 0
    return titles[i if 0 <= i < len(titles) else 0]


def write_answer(question, found, lang, seed):
    facts = found[4:].strip() if found.upper().startswith("YES:") else found
    return ask_long(f"Question: {question}\nANSWER FROM A SOURCE: {facts}\n\n"
                    f"Write this answer in {lang}, in one or two short sentences. "
                    "Keep the exact meaning.", seed, n=120).strip()


def agent(question, seed=0, max_tries=3):
    lang = detect_lang(question)
    tried = []
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    first = ask_long(LOOKUP_PROMPT + "Question: " + eng, seed, n=60).strip()
    m = re.search(r'lookup\("(.+?)"\)', first)
    if not m:
        if lang == "English" or len(first) < 200 and not first.lower().startswith(("i need", "could you")):
            return first
        return NOT_FOUND[lang]
    title = m.group(1)
    own = title
    if not exists(title):
        title = pick_from_search(eng, seed, exclude=tried, extra_query=own, size=6)
    for attempt in range(max_tries):
        if not title or title in tried:
            break
        tried.append(title)
        found = chunk_scan(title, eng)
        if found:
            return write_answer(question, found, lang, seed)
        title = pick_from_search(eng, seed, exclude=tried, extra_query=own, size=8)
    return NOT_FOUND[lang]



# ---- v7 (13 Oct 2026): list scan, current-fact scan, website crawl, chat layer, Roman input, brand spelling. Not tested standalone. ----
import time, json, urllib.parse, urllib.robotparser
from collections import deque
ZW = '\u200c'
BRAND = 'शत्\u200cचन्द्र शील्ड'
NAMES = {'Shatachandra Shield': 'शत्\u200cचन्द्र शील्ड'}
WRONG_NAME_FORMS = ['शत् चन्द्र शील्ड', 'शत् चन्द्र', 'शत् चंद्र', 'शत्चन्द्र', 'शत्चंद्र', 'शताचंद्र', 'शताचन्द्र', 'शतचंद्र', 'शतचन्द्र']
STEM = 'शत्\u200cचन्द्र'
STRONG = {'hai', 'tha', 'pehle', 'kahan', 'kisne', 'kaun', 'thi', 'aap', 'kab', 'shukriya', 'kitna', 'pehla', 'dhanyavad', 'mein', 'kitne', 'nahin', 'kaisa', 'nahi', 'rashtriya', 'namaste', 'dhanyawad', 'namaskar', 'hua', 'hui', 'kaise', 'hain', 'kya', 'janm'}
WEAK = {'ko', 'bharat', 'sa', 'ke', 'ka'}
TE = {'undi', 'unnadi', 'enduku', 'bharatadesha', 'sthapincharu', 'ela', 'evaru', 'emiti', 'edi', 'eppudu', 'enni', 'jateeya', 'dhanyavadalu', 'namaskaram', 'ekkada'}
SITES = {'auraworks': 'https://auraworks.online', 'shatachandra': 'https://www.shatachandrashield.xyz'}
CUR_WORDS = ['current', 'present', 'latest', 'right now', 'वर्तमान', 'अभी', 'ప్రస్తుత', 'ఇప్పుడు']
LIST_WORDS_EN = ['names of', 'list of', 'which are', 'what are the', 'name all', 'name the']
LIST_WORDS_HI = ['नाम', 'सूची', 'कौन-कौन', 'कौन कौन', 'naam', 'naam bataiye', 'list']
LIST_WORDS_TE = ['పేర్లు', 'జాబితా']
CLOSE = 'हमारी सुविधा प्राप्त करने के लिए धन्यवाद।\nआपकी सेवा करके हमें सुख की अनुभूति होती है।\nकृपया जब भी आपको हमारी आवश्यकता पड़े, आप हमें पुनः संदेश भेज सकते हैं।\nआपका दिन शुभ हो।'
RULES = {'Hindi': {'greet': ['नमस्ते', 'नमस्कार', 'हैलो', 'namaste', 'namaskar'], 'thanks': ['धन्यवाद', 'शुक्रिया', 'आभार', 'dhanyawad', 'dhanyavad', 'shukriya'], 'no': ['नहीं', 'बस', 'nahi', 'nahin', 'bas'], 'hello': 'नमस्ते। मेरे द्वारा आपकी क्या सहायता की जा सकती है?', 'welcome': 'इसकी कोई आवश्यकता नहीं।\nक्या मेरे द्वारा आपकी कोई और सहायता की जा सकती है?', 'close': 'हमारी सुविधा प्राप्त करने के लिए धन्यवाद।\nआपकी सेवा करके हमें सुख की अनुभूति होती है।\nकृपया जब भी आपको हमारी आवश्यकता पड़े, आप हमें पुनः संदेश भेज सकते हैं।\nआपका दिन शुभ हो।'}, 'English': {'greet': ['hello', 'hi', 'hey', 'namaste', 'namaskar'], 'thanks': ['thanks', 'thank you', 'dhanyawad', 'shukriya'], 'no': ['no', 'no thanks', "that's all", 'nothing else', 'bye'], 'hello': 'Namaste. How may I help you?', 'welcome': "You're welcome.\nIs there anything else I can help you with?", 'close': 'Thank you for using our service. Have a good day.'}, 'Telugu': {'greet': ['నమస్కారం', 'నమస్తే', 'హలో'], 'thanks': ['ధన్యవాదాలు', 'ధన్యవాదములు', 'థాంక్స్'], 'no': [], 'hello': '(Telugu greeting wording not set yet)', 'welcome': '(Telugu reply to thanks not set yet)', 'close': '(Telugu closing not set yet)'}}
_site_cache = {}

def detect_lang2(q):
    if re.search(r"[\u0C00-\u0C7F]", q): return "Telugu"
    if re.search(r"[\u0900-\u097F]", q): return "Hindi"
    w = re.findall(r"[a-z]+", q.lower())
    te = sum(x in TE for x in w)
    hi = 2 * sum(x in STRONG for x in w) + sum(x in WEAK for x in w)
    if te >= 1 and te * 2 > hi: return "Telugu"
    if hi >= 2: return "Hindi"
    return "English"

def pick_from_search(eng, seed, exclude=(), extra_query=None, size=6):
    titles = []
    for qq in [eng] + ([extra_query] if extra_query else []):
        for t in search_titles(qq, 8):
            if t not in exclude and t not in titles:
                titles.append(t)
    titles = titles[:size]
    if not titles:
        return None
    menu = "\n".join(f"{i+1}. {t}" for i, t in enumerate(titles))
    pick = ask_long("Question: " + eng + "\n\nWikipedia pages:\n" + menu +
                    "\n\nReply with only the number of the page most likely to contain the answer.",
                    seed, n=5).strip()
    m = re.search(r"\d", pick)
    i = int(m.group()) - 1 if m else 0
    return titles[i if 0 <= i < len(titles) else 0]

def write_answer(question, found, lang, seed):
    facts = found[4:].strip() if found.upper().startswith("YES:") else found
    if lang != "Hindi":
        return ask_long(f"Question: {question}\nANSWER FROM A SOURCE: {facts}\n\n"
                        f"Write this answer in {lang}, in one or two short sentences. "
                        "Keep the exact meaning.", seed, n=120).strip()
    pat = re.compile(r"shatachandra\s*shield", re.I)
    q2, f2 = pat.sub("NAMEX", question), pat.sub("NAMEX", facts)
    out = ask_long(f"Question: {q2}\nANSWER FROM A SOURCE: {f2}\n\n"
                   "Write this answer in Hindi using Devanagari script, in one or two short sentences. "
                   "Keep the exact meaning. The token NAMEX is a product name: keep NAMEX exactly as it is. "
                   "Keep technical terms such as ALLOW, DENY and ML-KEM in English letters.",
                   seed, n=120).strip()
    out = out.replace("NAMEX", BRAND)
    for w in WRONG_NAME_FORMS:
        out = out.replace(w, STEM)
    return out

def make_chunks(title, max_chars=4500):
    blocks = [b for b in page_blocks2(title) if not b.startswith("^")]
    chunks, cur = [], ""
    for b in blocks:
        if len(cur) + len(b) > max_chars:
            chunks.append(cur)
            cur = ""
        cur += b[:300] + "\n"
    if cur:
        chunks.append(cur)
    return chunks

def list_scan(title, eng, max_chunks=16):
    parts = []
    for c in make_chunks(title)[:max_chunks]:
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c +
                     "\n\nList only the items that answer the question and are stated in this TEXT, "
                     "separated by commas. If there are none, reply NONE.", 0, n=80).strip()
        if r and not r.upper().startswith("NONE"):
            parts.append(r)
    if not parts:
        return None
    return ask_long("Question: " + eng + "\n\nPARTIAL LISTS:\n" + "\n".join(parts) +
                    "\n\nCombine into one list without duplicates. Use only items from the partial lists. "
                    "Separate with commas.", 0, n=200).strip()

def current_scan(title, eng, max_chunks=16):
    parts = []
    for c in make_chunks(title)[:max_chunks]:
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c +
                     "\n\nIf this TEXT states an answer, reply with the answer and the year or date it applies to, "
                     "in one line. If not, reply NONE.", 0, n=60).strip()
        if r and not r.upper().startswith("NONE"):
            parts.append(r)
    if not parts:
        return None
    return ask_long("Question: " + eng + "\n\nCANDIDATE ANSWERS:\n" + "\n".join(parts) +
                    "\n\nChoose the most recent answer. Reply with the answer and the date or year it applies to.",
                    0, n=80).strip()

def wants_list(q, eng):
    e = eng.lower()
    return (any(w in e for w in LIST_WORDS_EN) or any(w in q.lower() for w in LIST_WORDS_HI + LIST_WORDS_TE)) and \
           any(c.isdigit() or w in e for c in "0123456789" for w in ["twelve","eleven","ten","nine","eight","seven","six","five","four","three","two","all","list"]) or \
           e.startswith(("list ", "name all", "what are the names"))

def page_edited(title):
    j = wiki_get({"action": "query", "format": "json", "formatversion": 2,
                  "prop": "revisions", "rvprop": "timestamp", "titles": title})
    return j["query"]["pages"][0]["revisions"][0]["timestamp"]

def fetch_blocks(url):
    r = requests.get(url, headers=H, timeout=20)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    lines = [re.sub(r"\s+", " ", x).strip() for x in soup.get_text("\n").split("\n")]
    blocks = [b for b in lines if len(b) > 1]
    links = [urllib.parse.urljoin(url, a["href"]).split("#")[0] for a in soup.find_all("a", href=True)]
    return blocks, links

def crawl(start, max_pages=6):
    host = urllib.parse.urlparse(start).netloc
    rp = urllib.robotparser.RobotFileParser()
    try:
        rp.set_url(urllib.parse.urljoin(start, "/robots.txt"))
        rp.read()
    except Exception:
        rp = None
    seen, queue, pages = set(), deque([start]), {}
    while queue and len(pages) < max_pages:
        u = queue.popleft()
        if u in seen:
            continue
        seen.add(u)
        if rp and not rp.can_fetch("*", u):
            continue
        try:
            blocks, links = fetch_blocks(u)
        except Exception:
            continue
        pages[u] = blocks
        for l in links:
            if (urllib.parse.urlparse(l).netloc == host and l not in seen
                    and not l.lower().endswith((".png", ".jpg", ".pdf", ".zip", ".mp4"))):
                queue.append(l)
        time.sleep(1)
    return pages

def site_chunks(pages, max_chars=4500):
    out = []
    for u, blocks in pages.items():
        cur = ""
        for b in blocks:
            if len(cur) + len(b) > max_chars:
                out.append((u, cur))
                cur = ""
            cur += b[:300] + "\n"
        if cur:
            out.append((u, cur))
    return out

def site_answer(question, eng, name, lang, seed=0):
    if name not in _site_cache:
        _site_cache[name] = (crawl(SITES[name], max_pages=25), time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime()))
    pages, read_at = _site_cache[name]
    words = [w for w in toks(eng) if len(w) > 2]
    chunks = sorted(site_chunks(pages), key=lambda uc: -sum(w in uc[1].lower() for w in words))
    for u, c in chunks[:20]:
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c +
                     "\n\nDoes this TEXT directly state the answer? If yes, reply 'YES: ' followed by the answer. "
                     "If not, reply 'NO'.", seed, n=60).strip()
        if r.upper().startswith("YES"):
            return write_answer(question, r, lang, seed) + f"\n(Source: {u}, read {read_at})"
    return NOT_FOUND[lang]

def chat_turn(msg, st):
    lang = detect_lang(msg)
    R = RULES[lang]
    m = msg.strip().lower().strip("!.,?। ")
    if st["asked"] and any(m in RULES[l]["no"] for l in RULES):
        st["asked"] = False
        return RULES[st.get("lang", lang)]["close"]
    if m in R["greet"]:
        st["first"], st["asked"] = False, False
        return R["hello"]
    if any(w in m for w in R["thanks"]):
        st["asked"], st["lang"] = True, lang
        return R["welcome"]
    ans = agent(msg, 0)
    out = ("नमस्ते।\n" if st["first"] and lang == "Hindi" else "") + ans
    st["first"], st["asked"] = False, False
    return out

def run_chat(msgs):
    st = {"first": True, "asked": False}
    for m in msgs:
        print("USER :", m)
        print("AGENT:", chat_turn(m, st).replace("\n", " | "))
    print()

detect_lang = detect_lang2
_agent_single = agent
def agent(question, seed=0, max_tries=3):
    lang = detect_lang(question)
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    if not wants_list(question, eng):
        return _agent_single(question, seed, max_tries)
    first = ask_long(LOOKUP_PROMPT + "Question: " + eng, seed, n=60).strip()
    m = re.search(r'lookup\("(.+?)"\)', first)
    if not m:
        return _agent_single(question, seed, max_tries)
    title = m.group(1)
    if not exists(title):
        title = pick_from_search(eng, seed, extra_query=title, size=6)
    found = list_scan(title, eng) if title else None
    if not found:
        return NOT_FOUND[lang]
    return write_answer(question, found, lang, seed)

_agent_prev = agent
def agent(question, seed=0, max_tries=3):
    lang = detect_lang(question)
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    if not any(w in (question + " " + eng).lower() for w in CUR_WORDS):
        return _agent_prev(question, seed, max_tries)
    first = ask_long(LOOKUP_PROMPT + "Question: " + eng, seed, n=60).strip()
    m = re.search(r'lookup\("(.+?)"\)', first)
    if not m:
        return _agent_prev(question, seed, max_tries)
    title = m.group(1)
    if not exists(title):
        title = pick_from_search(eng, seed, extra_query=title, size=6)
    found = current_scan(title, eng) if title else None
    if not found:
        return NOT_FOUND[lang]
    try:
        edited = page_edited(title)[:10]
    except Exception:
        edited = "unknown"
    return write_answer(question, found, lang, seed) + f"\n(Source: Wikipedia, {title}, last edited {edited})"

_agent_nosite = agent
def agent(question, seed=0, max_tries=3):
    low = question.lower()
    for name in SITES:
        if name in low or name.replace(" ", "") in low.replace(" ", ""):
            lang = detect_lang(question)
            eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                           seed, n=40).strip()
            return site_answer(question, eng, name, lang, seed)
    return _agent_nosite(question, seed, max_tries)



# ======== v8 patches (14 Oct 2026, Colab T4). Run AFTER the model is loaded. ========
import contextlib, time
if not hasattr(model, "disable_adapter"):
    model.disable_adapter = lambda: contextlib.nullcontext()

# 1) source line on every answer
LAST = {"title": None}
def _wrap_scan(fn):
    def w(title, *a, **k):
        r = fn(title, *a, **k)
        if r:
            LAST["title"] = title
        return r
    return w
chunk_scan = _wrap_scan(chunk_scan)
list_scan = _wrap_scan(list_scan)
current_scan = _wrap_scan(current_scan)

_v7_agent = agent
def agent(question, seed=0, max_tries=3):
    LAST["title"] = None
    out = _v7_agent(question, seed, max_tries)
    if "(Source:" in out or out in NOT_FOUND.values():
        return out
    if LAST["title"]:
        try:
            ed = page_edited(LAST["title"])[:10]
        except Exception:
            ed = "unknown"
        return out + f"\n(Source: Wikipedia, {LAST['title']}, last edited {ed})"
    return out + "\n(No source checked: answer from the model's memory)"

# 2) protected names in Hindi / Telugu / Roman -> official sites
SITES["nirbhedya"] = "https://nirbhedya.in"
NAME_FORMS = {
 "shatachandra": ["शतचन्द्र","शतचंद्र","शत्चन्द्र","शत्चंद्र","शताचंद्र","शताचन्द्र","శతచంద్ర","shatachandra","shatchandra","shatachandrashield"],
 "nirbhedya": ["निर्भेद्य","నిర్భేద్య","nirbhedya","nirbhedhya"],
 "auraworks": ["औरावर्क्स","ऑरावर्क्स","ఔరావర్క్స్","auraworks"],
}
ENGLISH = {"shatachandra": "Shatachandra Shield", "nirbhedya": "Nirbhedya", "auraworks": "AuraWorks"}
MYTH = ["myth", "पौराणिक", "विष्णु", "पुराण", "పురాణ"]

def _norm(s):
    return "".join(ch for ch in s.lower() if ch not in " \u200c\u200d")

_agent_before_names = agent
def agent(question, seed=0, max_tries=3):
    n = _norm(question)
    if not any(m in n for m in MYTH):
        for key, forms in NAME_FORMS.items():
            if any(_norm(f) in n for f in forms):
                lang = detect_lang(question)
                q2 = question
                for f in forms:
                    q2 = q2.replace(f, ENGLISH[key])
                eng = ask_long("Translate this question to English. Output only the translation.\n\n" + q2,
                               seed, n=40).strip()
                return site_answer(question, eng, key, lang, seed)
    return _agent_before_names(question, seed, max_tries)

# 3) final Hindi spellings (author, 14 Oct 2026)
GOOD_SH, GOOD_AW = "शतचन्द्र", "औरा वर्क्स"
BAD_SH = ["शत्\u200cचन्द्र","शत्\u200cचंद्र","शत् चन्द्र","शत् चंद्र","शत्चन्द्र","शत्चंद्र","शतचंद्र","शताचंद्र","शताचन्द्र"]
BAD_AW = ["ऑरा वर्क्स","ऑरावर्क्स","औरावर्क्स"]
BRAND = GOOD_SH + " शील्ड"
STEM = GOOD_SH
NAMES = {"Shatachandra Shield": BRAND}
WRONG_NAME_FORMS = BAD_SH

def fix_names(text):
    for b in BAD_SH:
        text = text.replace(b, GOOD_SH)
    for b in BAD_AW:
        text = text.replace(b, GOOD_AW)
    return text

_agent_before_fix = agent
def agent(question, seed=0, max_tries=3):
    return fix_names(_agent_before_fix(question, seed, max_tries))

# 4) more room for Telugu / Hindi answers (no cut-off)
def write_answer(question, found, lang, seed):
    facts = found[4:].strip() if found.upper().startswith("YES:") else found
    n = {"Telugu": 400, "Hindi": 250}.get(lang, 150)
    return ask_long(f"Question: {question}\nANSWER FROM A SOURCE: {facts}\n\n"
                    f"Write this answer in {lang}, in one or two short sentences. "
                    "Keep the exact meaning.", seed, n=n).strip()


# ======== v9 patches (14 Oct 2026): force lookup + ask-back for ambiguous names ========
import requests as _rq

# 1) never answer a fact question from memory: search Wikipedia and read the page
_agent_before_force = agent
def agent(question, seed=0, max_tries=3):
    out = _agent_before_force(question, seed, max_tries)
    if "(No source checked" not in out:
        return out
    lang = detect_lang(question)
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    title = pick_from_search(eng, seed, size=6)
    if not title:
        return out
    scan = list_scan if wants_list(question, eng) else chunk_scan
    found = scan(title, eng)
    if not found:
        return out
    try:
        ed = page_edited(title)[:10]
    except Exception:
        ed = "unknown"
    return fix_names(write_answer(question, found, lang, seed)) + f"\n(Source: Wikipedia, {title}, last edited {ed})"

# 2) ask back when a bare topic is a Wikipedia disambiguation page
_UA = {"User-Agent": "ShatachandraAI-Lab/0.1 (research; github.com/Vivek76760/shatachandra-ai-lab-open)"}
ASK = {"English": "This name has several meanings. Which one do you mean?",
       "Hindi": "इस नाम के कई अर्थ हैं। आप किसके बारे में पूछ रहे हैं?",
       "Telugu": "ఈ పేరుకు అనేక అర్థాలు ఉన్నాయి. మీరు ఏది అడుగుతున్నారు?"}

def disambig_options(title, n=12):
    q = {"action": "query", "titles": title, "redirects": 1, "format": "json",
         "prop": "pageprops|links", "ppprop": "disambiguation", "pllimit": "max", "plnamespace": 0}
    d = None
    for _try in range(2):
        try:
            d = _rq.get(API, params=q, headers=_UA, timeout=20).json()
            break
        except Exception:
            time.sleep(2)
    if d is None:
        return None, None
    for pg in d.get("query", {}).get("pages", {}).values():
        if "disambiguation" in pg.get("pageprops", {}):
            links = [l["title"] for l in pg.get("links", [])]
            t = title.lower().replace(" (disambiguation)", "")
            gen = [x for x in links if x.lower().startswith(t + " (") and x[len(t) + 2:len(t) + 3].islower()]
            spec = [x for x in links if x.lower().startswith(t) and x not in gen]
            rest = [x for x in links if x not in gen and x not in spec]
            return (gen + spec + rest)[:n], pg["title"]
    return None, None

_agent_before_amb = agent
def agent(question, seed=0, max_tries=3):
    lang = detect_lang(question)
    base = question if lang == "English" else ask_long(
        "Translate this question to English. Output only the translation.\n\n" + question, seed, n=40).strip()
    topic = re.sub(r"^(please\s+)?(tell me about|what is|what are|who is|who was|explain|about|define)\s+", "",
                   base.strip().rstrip("?.! "), flags=re.I)
    topic = re.sub(r"^(a|an|the)\s+", "", topic, flags=re.I).strip()
    known = any(_norm(f) in _norm(question) or _norm(f) in _norm(topic) for fs in NAME_FORMS.values() for f in fs)
    if topic and len(topic.split()) <= 2 and not known:
        opts, real = disambig_options(topic.title())
        if opts:
            return ASK[lang] + "\n- " + "\n- ".join(opts) + f"\n(Source: Wikipedia disambiguation page, {real})"
    return _agent_before_amb(question, seed, max_tries)

# 3) show protected names in the answer's own script (Hindi / Telugu), source line untouched
NAME_LOCAL = {
 "Hindi": {"Shatachandra Shield": "शतचन्द्र शील्ड", "Aura Works": "औरा वर्क्स", "AuraWorks": "औरा वर्क्स",
           "Nirbhedya": "निर्भेद्य", "Shatachandra": "शतचन्द्र"},
 "Telugu": {"Shatachandra Shield": "శతచంద్ర షీల్డ్", "Aura Works": "ఔరా వర్క్స్", "AuraWorks": "ఔరా వర్క్స్",
            "Nirbhedya": "నిర్భేద్య", "Shatachandra": "శతచంద్ర"},
}
def localize_names(text, lang):
    m = NAME_LOCAL.get(lang)
    if not m:
        return text
    body, sep, src = text.partition("\n(Source:")
    for k, v in m.items():
        body = body.replace(k, v)
    return body + sep + src

_agent_before_local = agent
def agent(question, seed=0, max_tries=3):
    return localize_names(_agent_before_local(question, seed, max_tries), detect_lang(question))


# ======== v10 patches (14 Oct 2026): speed + correct edit date ========
# 1) translate each question only once (v7 layers translated it 3 times)
_ask_long_orig = ask_long
_TR = {}
def ask_long(*a, **k):
    if a and isinstance(a[0], str) and a[0].startswith("Translate this question to English"):
        key = (a, tuple(sorted(k.items())))
        if key not in _TR:
            if len(_TR) > 500:
                _TR.clear()
            _TR[key] = _ask_long_orig(*a, **k)
        return _TR[key]
    return _ask_long_orig(*a, **k)

# 2) English: the page scan already gives a full sentence, skip the rewrite (saves about 4 s)
_write_answer_full = write_answer
def write_answer(question, found, lang, seed):
    if lang == "English":
        facts = found[4:].strip() if found.upper().startswith("YES:") else found
        if len(facts.split()) >= 5:
            return facts
    return _write_answer_full(question, found, lang, seed)

# 3) edit date of the real page (follow redirects)
def page_edited(title):
    d = _rq.get(API, params={"action": "query", "titles": title, "prop": "revisions",
                             "rvprop": "timestamp", "redirects": 1, "format": "json"},
                headers=_UA, timeout=20).json()
    for pg in d.get("query", {}).get("pages", {}).values():
        revs = pg.get("revisions")
        if revs:
            return revs[0]["timestamp"]
    return "unknown"


# ======== v11 patches (14 Oct 2026): clean script check for Hindi / Telugu answers ========
import unicodedata
_ALLOWED = {"Telugu": ("TELUGU", "LATIN"), "Hindi": ("DEVANAGARI", "LATIN")}

def _foreign(text, lang):
    ok = _ALLOWED.get(lang)
    if not ok:
        return False
    for ch in text:
        if ch.isalpha() and not unicodedata.name(ch, "").startswith(ok):
            return True
    return False

_write_answer_prev = write_answer
def write_answer(question, found, lang, seed):
    out = _write_answer_prev(question, found, lang, seed)
    if not _foreign(out, lang):
        return out
    facts = found[4:].strip() if found.upper().startswith("YES:") else found
    n = {"Telugu": 400, "Hindi": 250}.get(lang, 150)
    for k in (0, 1):
        out2 = ask_long(f"Question: {question}\nANSWER FROM A SOURCE: {facts}\n\n"
                        f"Write this answer in {lang}, in one or two short sentences. Keep the exact meaning. "
                        f"Use only {lang} script, but write all place names and person names in English letters.",
                        seed + k, n=n).strip()
        if not _foreign(out2, lang):
            return out2
    return facts + f"\n(I could not write this cleanly in {lang}; this is the source text in English.)"


# ======== v12 patches (14 Oct 2026): web search fallback (ddgs) with a safe page reader ========
import socket, ipaddress
from ddgs import DDGS

_BLOCK_HOSTS = ("youtube.com", "facebook.com", "instagram.com", "twitter.com", "x.com",
                "pinterest.com", "quora.com", "tiktok.com")

def _public_url(u):
    try:
        p = urllib.parse.urlparse(u)
        if p.scheme not in ("http", "https") or not p.hostname:
            return False
        h = p.hostname.lower()
        if any(h == b or h.endswith("." + b) for b in _BLOCK_HOSTS):
            return False
        for info in socket.getaddrinfo(h, None):
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
                return False
        return True
    except Exception:
        return False

def safe_blocks(url, max_bytes=1500000):
    if not _public_url(url):
        return None
    p = urllib.parse.urlparse(url)
    try:
        rr = requests.get(f"{p.scheme}://{p.netloc}/robots.txt", headers=H, timeout=10)
        if rr.status_code == 200:
            rp = urllib.robotparser.RobotFileParser()
            rp.parse(rr.text.splitlines())
            if not rp.can_fetch("*", url):
                return None
    except Exception:
        pass
    r = requests.get(url, headers=H, timeout=15, stream=True)
    if r.status_code != 200 or "html" not in r.headers.get("content-type", "").lower():
        return None
    if not _public_url(r.url):
        return None
    data = r.raw.read(max_bytes, decode_content=True)
    soup = BeautifulSoup(data.decode(r.encoding or "utf-8", errors="ignore"), "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    lines_ = [re.sub(r"\s+", " ", x).strip() for x in soup.get_text("\n").split("\n")]
    return [b for b in lines_ if len(b) > 1]

def web_answer(question, eng, lang, seed=0, max_pages=3, max_scans=6):
    try:
        res = DDGS().text(eng, max_results=8)
    except Exception:
        return None
    words = [w for w in toks(eng) if len(w) > 2]
    cands, used = [], 0
    for r in res:
        if used >= max_pages:
            break
        u = r.get("href", "")
        try:
            blocks = safe_blocks(u)
        except Exception:
            blocks = None
        if not blocks:
            continue
        used += 1
        cands += site_chunks({u: blocks})
    cands.sort(key=lambda uc: -sum(w in uc[1].lower() for w in words))
    read_at = time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime())
    for u, c in cands[:max_scans]:
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c +
                     "\n\nDoes this TEXT directly state the answer? If yes, reply 'YES: ' followed by the answer. "
                     "If not, reply 'NO'.", seed, n=60).strip()
        if r.upper().startswith("YES"):
            return write_answer(question, r, lang, seed) + f"\n(Source: {u}, read {read_at})"
    return None

_agent_before_web = agent
def agent(question, seed=0, max_tries=3):
    out = _agent_before_web(question, seed, max_tries)
    if out not in NOT_FOUND.values():
        return out
    n = _norm(question)
    if any(_norm(f) in n for fs in NAME_FORMS.values() for f in fs):
        return out
    lang = detect_lang(question)
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    got = web_answer(question, eng, lang, seed)
    return got if got else out


# ======== v13 (14 Oct 2026): cross-check web answers, official sites first ========
OFFICIAL = (".gov.in", ".nic.in", ".gov", ".edu", ".ac.in", ".edu.in")
DISAGREE = {"English": "The sources I found disagree. Please check:",
            "Hindi": "मुझे मिले स्रोतों में अलग-अलग जानकारी है। कृपया जाँच लें:",
            "Telugu": "నాకు దొరికిన మూలాల్లో వేర్వేరు సమాచారం ఉంది. దయచేసి సరిచూడండి:"}
_SCAN_TAIL = "\n\nDoes this TEXT directly state the answer? If yes, reply 'YES: ' followed by the answer. If not, reply 'NO'."

def _official(u):
    return (urllib.parse.urlparse(u).hostname or "").endswith(OFFICIAL)

def _nums(s):
    return set(re.findall(r"\d+(?:\.\d+)?", s.replace(",", "")))

def web_answer(question, eng, lang, seed=0, max_pages=4, max_scans=8):
    try:
        res = DDGS().text(eng, max_results=10)
    except Exception:
        return None
    words = [w for w in toks(eng) if len(w) > 2]
    cands, used = [], 0
    for r in res:
        if used >= max_pages:
            break
        u = r.get("href", "")
        try:
            blocks = safe_blocks(u)
        except Exception:
            blocks = None
        if not blocks:
            continue
        used += 1
        cands += site_chunks({u: blocks})
    cands.sort(key=lambda uc: (not _official(uc[0]), -sum(w in uc[1].lower() for w in words)))
    hits, scans = [], 0
    for u, c in cands:
        if scans >= max_scans or len(hits) >= 3:
            break
        if any(u == h[0] for h in hits):
            continue
        scans += 1
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c + _SCAN_TAIL, seed, n=60).strip()
        if r.upper().startswith("YES"):
            hits.append((u, r))
    if not hits:
        return None
    rd = time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime())
    if len(hits) == 1:
        return (write_answer(question, hits[0][1], lang, seed) + f"\n(Source: {hits[0][0]}, read {rd})"
                + "\n(Only one source found, not cross-checked.)")
    base = _nums(hits[0][1])
    if all((not base) or (not _nums(h[1])) or _nums(h[1]) == base for h in hits):
        return write_answer(question, hits[0][1], lang, seed) + "".join(f"\n(Source: {u}, read {rd})" for u, _ in hits)
    out = [DISAGREE.get(lang, DISAGREE["English"])]
    out += [f"- {write_answer(question, r, lang, seed)} (Source: {u})" for u, r in hits]
    return "\n".join(out) + f"\n(read {rd})"

_agent_before_web2 = agent
def agent(question, seed=0, max_tries=3):
    out = _agent_before_web2(question, seed, max_tries)
    if "(No source checked" not in out:
        return out
    n = _norm(question)
    if any(_norm(f) in n for fs in NAME_FORMS.values() for f in fs):
        return out
    lang = detect_lang(question)
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    return web_answer(question, eng, lang, seed) or out
# END


# ======== v14 (14 Oct 2026): read files (txt, md, json, csv, html, pdf) and answer from them ========
import os, json as _json, csv as _csv

def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".md", ".py", ".log"):
        return open(path, encoding="utf-8", errors="ignore").read()
    if ext == ".json":
        return _json.dumps(_json.load(open(path, encoding="utf-8")), ensure_ascii=False, indent=1)
    if ext == ".csv":
        return "\n".join(", ".join(row) for row in _csv.reader(open(path, encoding="utf-8", errors="ignore")))
    if ext in (".html", ".htm"):
        soup = BeautifulSoup(open(path, encoding="utf-8", errors="ignore").read(), "html.parser")
        for t in soup(["script", "style", "noscript"]):
            t.decompose()
        return soup.get_text("\n")
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    raise ValueError("unsupported file type: " + ext)

def file_answer(question, path, seed=0, max_scans=10):
    lang = detect_lang(question)
    eng = ask_long("Translate this question to English. Output only the translation.\n\n" + question,
                   seed, n=40).strip()
    blocks = [re.sub(r"\s+", " ", b).strip() for b in read_any(path).split("\n")]
    blocks = [b for b in blocks if len(b) > 1]
    words = [w for w in toks(eng) if len(w) > 2]
    chunks = site_chunks({path: blocks})
    chunks.sort(key=lambda uc: -sum(w in uc[1].lower() for w in words))
    for u, c in chunks[:max_scans]:
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c + _SCAN_TAIL, seed, n=60).strip()
        if r.upper().startswith("YES"):
            return write_answer(question, r, lang, seed) + f"\n(Source: file {os.path.basename(path)})"
    return NOT_FOUND[lang]
# END


# ======== v15 (14 Oct 2026): write Python, run it, read the result, fix errors (max 3 tries) ========
import subprocess, sys, tempfile
CODE_PROMPT = ("Write a complete Python 3 program for this task. Output ONLY the code. "
               "The program must print its final result. Task: ")

def _extract_code(t):
    m = re.search(r"`{3}(?:python)?\n(.*?)`{3}", t, re.S)
    return (m.group(1) if m else t).strip()

def run_python(code, timeout=30):
    d = tempfile.mkdtemp()
    p = os.path.join(d, "task.py")
    open(p, "w", encoding="utf-8").write(code)
    try:
        r = subprocess.run([sys.executable, p], capture_output=True, text=True, timeout=timeout, cwd=d)
        return r.returncode, (r.stdout + r.stderr)[-1500:], d
    except subprocess.TimeoutExpired:
        return 124, "timed out after %s seconds" % timeout, d

def code_task(task, seed=0, tries=3):
    prompt = CODE_PROMPT + task
    code, out, d = "", "", ""
    for i in range(tries):
        code = _extract_code(ask_long(prompt, seed + i, n=400))
        rc, out, d = run_python(code)
        if rc == 0 and out.strip():
            return {"ok": True, "tries": i + 1, "code": code, "output": out.strip(), "dir": d}
        prompt = (CODE_PROMPT + task + "\n\nYour previous code failed or printed nothing. Output:\n" + out +
                  "\n\nPrevious code:\n" + code + "\n\nFix it. Output ONLY the corrected code.")
    return {"ok": False, "tries": tries, "code": code, "output": out.strip(), "dir": d}
# END


# ======== v16 (14 Oct 2026): one agent chooses its tool (file reader, code runner, or normal lookup) ========
_FILE_RE = re.compile(r"[\w\-./]+\.(?:pdf|txt|md|json|csv|html?)\b", re.I)
CODE_WORDS = ("write a program", "write a python", "python program", "python script", "write code",
              "write a function", "write a script", "प्रोग्राम लिख", "प्रोग्राम बना", "कोड लिख",
              "कोड बना", "प्रिंट करो", "ప్రోగ్రామ్ రాయ", "కోడ్ రాయ")

_agent_before_route = agent
def agent(question, seed=0, max_tries=3):
    m = _FILE_RE.search(question)
    if m and os.path.exists(m.group(0)):
        return file_answer(question, m.group(0), seed)
    low = question.lower()
    if any(w in low for w in CODE_WORDS):
        r = code_task(question, seed)
        if r["ok"]:
            return ("Result:\n" + r["output"] + "\n\nCode used:\n" + r["code"] +
                    f"\n(Ran the code: no errors, tries: {r['tries']}. Running without errors does not prove "
                    "the result is what you wanted.)")
        return "I could not make a working program in 3 tries. Last output:\n" + r["output"]
    return _agent_before_route(question, seed, max_tries)
# END


# ======== v17 (14 Oct 2026): safer code extraction, longer code allowed, shorter programs asked for ========
CODE_PROMPT = ("Write a complete Python 3 program for this task. Keep it short: no comments, no docstrings. "
               "Output ONLY the code. The program must print its final result. Task: ")

def _extract_code(t):
    m = re.search(r"`{3}(?:python)?\s*\n(.*?)`{3}", t, re.S)
    if m:
        return m.group(1).strip()
    t = re.sub(r"^\s*`{3}(?:python)?\s*\n", "", t)
    return t.replace("`" * 3, "").strip()

def code_task(task, seed=0, tries=3):
    prompt = CODE_PROMPT + task
    code, out, d = "", "", ""
    for i in range(tries):
        code = _extract_code(ask_long(prompt, seed + i, n=900))
        rc, out, d = run_python(code)
        if rc == 0 and out.strip():
            return {"ok": True, "tries": i + 1, "code": code, "output": out.strip(), "dir": d}
        prompt = (CODE_PROMPT + task + "\n\nYour previous code failed or printed nothing. Output:\n" + out +
                  "\n\nPrevious code:\n" + code + "\n\nFix it. Output ONLY the corrected code.")
    return {"ok": False, "tries": tries, "code": code, "output": out.strip(), "dir": d}
# END


# ======== v18 (14 Oct 2026): edit files safely (work on a copy, save new file in 'outputs'), read docx/xlsx ========
import shutil

def run_python(code, timeout=30, files=()):
    d = tempfile.mkdtemp()
    for f in files:
        shutil.copy(f, d)
    p = os.path.join(d, "task.py")
    open(p, "w", encoding="utf-8").write(code)
    try:
        r = subprocess.run([sys.executable, p], capture_output=True, text=True, timeout=timeout, cwd=d)
        return r.returncode, (r.stdout + r.stderr)[-1500:], d
    except subprocess.TimeoutExpired:
        return 124, "timed out after %s seconds" % timeout, d

def code_task(task, seed=0, tries=3, files=()):
    names = [os.path.basename(f) for f in files]
    extra = ""
    if files:
        extra = (" These files are in the current folder: " + ", ".join(names) +
                 ". Never change the original: save the changes as a NEW file with a new name.")
    prompt = CODE_PROMPT + task + extra
    code, out, d = "", "", ""
    for i in range(tries):
        code = _extract_code(ask_long(prompt, seed + i, n=900))
        rc, out, d = run_python(code, files=files)
        if rc == 0 and out.strip():
            os.makedirs("outputs", exist_ok=True)
            made = []
            for fn in sorted(os.listdir(d)):
                src = os.path.join(d, fn)
                if fn == "task.py" or fn.startswith("__") or not os.path.isfile(src):
                    continue
                if fn in names and open(src, "rb").read() == open(files[names.index(fn)], "rb").read():
                    continue
                new = ("edited_" + fn) if fn in names else fn
                shutil.copy(src, os.path.join("outputs", new))
                made.append(new)
            return {"ok": True, "tries": i + 1, "code": code, "output": out.strip(), "dir": d, "files": made}
        prompt = (CODE_PROMPT + task + extra + "\n\nYour previous code failed or printed nothing. Output:\n" + out +
                  "\n\nPrevious code:\n" + code + "\n\nFix it. Output ONLY the corrected code.")
    return {"ok": False, "tries": tries, "code": code, "output": out.strip(), "dir": d, "files": []}

_read_any_prev = read_any
def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        import docx
        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path)
        return "\n".join(" | ".join("" if c is None else str(c) for c in row)
                         for ws in wb for row in ws.iter_rows(values_only=True))
    return _read_any_prev(path)

_FILE_RE = re.compile(r"[\w\-./]+\.(?:pdf|txt|md|json|csv|html?|docx|xlsx)\b", re.I)
_EDIT_WORDS = ("change ", "edit ", "replace ", "update ", "add ", "delete ", "remove ", "rename ")

_agent_before_edit = agent
def agent(question, seed=0, max_tries=3):
    m = _FILE_RE.search(question)
    if m and os.path.exists(m.group(0)) and any(w in question.lower() for w in _EDIT_WORDS):
        r = code_task("Write a Python program for this: " + question, seed, files=[m.group(0)])
        if r["ok"]:
            made = ("\n\nNew file(s) saved in the 'outputs' folder: " + ", ".join(r["files"])) if r["files"] else ""
            return ("Result:\n" + r["output"] + made +
                    "\n(The original file was not changed. Check the new file before you use it.)")
        return "I could not finish the file edit in 3 tries. Last output:\n" + r["output"]
    return _agent_before_edit(question, seed, max_tries)
# END


# ======== v19 (14 Oct 2026): a file edit counts as done when a new file is saved, even if nothing is printed ========
def _collect(d, names, files):
    found = []
    for fn in sorted(os.listdir(d)):
        src = os.path.join(d, fn)
        if fn == "task.py" or fn.startswith("__") or not os.path.isfile(src):
            continue
        if fn in names and open(src, "rb").read() == open(files[names.index(fn)], "rb").read():
            continue
        found.append((src, ("edited_" + fn) if fn in names else fn))
    return found

def code_task(task, seed=0, tries=3, files=()):
    names = [os.path.basename(f) for f in files]
    extra = ""
    if files:
        extra = (" These files are in the current folder: " + ", ".join(names) +
                 ". Never change the original: save the changes as a NEW file with a new name, "
                 "then print one short line saying which new file you saved.")
    prompt = CODE_PROMPT + task + extra
    code, out, d = "", "", ""
    for i in range(tries):
        code = _extract_code(ask_long(prompt, seed + i, n=900))
        rc, out, d = run_python(code, files=files)
        found = _collect(d, names, files) if rc == 0 else []
        if rc == 0 and (out.strip() or found):
            os.makedirs("outputs", exist_ok=True)
            made = []
            for src, new in found:
                shutil.copy(src, os.path.join("outputs", new))
                made.append(new)
            return {"ok": True, "tries": i + 1, "code": code, "output": out.strip() or "Done.", "dir": d, "files": made}
        prompt = (CODE_PROMPT + task + extra + "\n\nYour previous code failed, printed nothing, or saved no new file. "
                  "Output:\n" + out + "\n\nPrevious code:\n" + code + "\n\nFix it. Output ONLY the corrected code.")
    return {"ok": False, "tries": tries, "code": code, "output": out.strip(), "dir": d, "files": []}
# END


# ======== v20 (14 Oct 2026): file edits are compared with the original and self-checked ========
import difflib
from itertools import zip_longest

def file_diff(orig, new):
    ext = os.path.splitext(new)[1].lower()
    if ext == ".xlsx":
        import openpyxl
        a, b = openpyxl.load_workbook(orig), openpyxl.load_workbook(new)
        out = []
        for ws in b:
            wa = a[ws.title] if ws.title in a.sheetnames else None
            cells = {c.coordinate for row in ws.iter_rows() for c in row}
            if wa is not None:
                cells |= {c.coordinate for row in wa.iter_rows() for c in row}
            for co in sorted(cells):
                va, vb = (wa[co].value if wa is not None else None), ws[co].value
                if va != vb:
                    out.append(f"{ws.title}!{co}: {va!r} -> {vb!r}")
        return out[:30]
    if ext == ".docx":
        import docx
        pa = [p.text for p in docx.Document(orig).paragraphs]
        pb = [p.text for p in docx.Document(new).paragraphs]
        return [f"paragraph {i + 1}: {x!r} -> {y!r}" for i, (x, y) in enumerate(zip_longest(pa, pb)) if x != y][:30]
    la, lb = read_any(orig).splitlines(), read_any(new).splitlines()
    return [l for l in difflib.unified_diff(la, lb, lineterm="", n=0) if not l.startswith(("---", "+++", "@@"))][:30]

_agent_before_verify = agent
def agent(question, seed=0, max_tries=3):
    m = _FILE_RE.search(question)
    if not (m and os.path.exists(m.group(0)) and any(w in question.lower() for w in _EDIT_WORDS)):
        return _agent_before_verify(question, seed, max_tries)
    path, note, dtext, v = m.group(0), "", "", ""
    for rnd in range(2):
        r = code_task("Write a Python program for this: " + question + note, seed + rnd * 3, files=[path])
        if not (r["ok"] and r["files"]):
            return "I could not finish the file edit. Last output:\n" + r["output"]
        newp = os.path.join("outputs", r["files"][0])
        try:
            diff = file_diff(path, newp)
        except Exception as e:
            diff = ["(could not compare: %s)" % e]
        dtext = "\n".join(diff) or "(no change found)"
        v = ask_long(f"Request: {question}\n\nChanges made to the file:\n{dtext}\n\n"
                     "Does this change exactly match the request, with nothing else changed? "
                     "Reply 'YES' or 'NO: reason'.", seed, n=60).strip()
        if v.upper().startswith("YES"):
            return ("Result:\n" + r["output"] + "\n\nWhat changed:\n" + dtext +
                    f"\n\nNew file saved in the 'outputs' folder: {r['files'][0]}\n"
                    "(Self-check: the change matches the request. The original file was not changed.)")
        os.remove(newp)
        note = (f"\n\nA previous attempt was wrong. It changed:\n{dtext}\nProblem: {v}\nDo it correctly.")
    return ("I could not make this edit correctly in 2 rounds. The last attempt changed:\n" + dtext +
            "\n(Self-check said: " + v + ")")
# END


# ======== v21 (14 Oct 2026): show the model the file content before editing; honest wording about the self-check ========
def describe_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        import openpyxl
        lines = []
        for ws in openpyxl.load_workbook(path):
            lines.append(f"Sheet '{ws.title}':")
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        lines.append(f"{c.coordinate} = {c.value!r}")
        return "\n".join(lines[:80])
    if ext == ".docx":
        import docx
        return "\n".join(f"paragraph {i + 1}: {p.text}" for i, p in enumerate(docx.Document(path).paragraphs))[:1500]
    return read_any(path)[:1500]

_code_task_prev = code_task
def code_task(task, seed=0, tries=3, files=()):
    if files:
        try:
            task = task + "\n\nCurrent content of " + os.path.basename(files[0]) + ":\n" + describe_file(files[0])
        except Exception:
            pass
    return _code_task_prev(task, seed, tries, files)

_agent_before_wording = agent
def agent(question, seed=0, max_tries=3):
    out = _agent_before_wording(question, seed, max_tries)
    return out.replace("(Self-check: the change matches the request. The original file was not changed.)",
                       "(The model's own check said this matches the request, but that check is not reliable: "
                       "please read 'What changed' yourself. The original file was not changed.)")
# END


# ======== v22 (14 Oct 2026): glossary of local terms swapped for the right English term before translating ========
GLOSSARY = {
    "దండి యాత్ర": "Dandi March",
    "దండి సత్యాగ్రహం": "Dandi March",
    "దండి": "Dandi",
    "दांडी मार्च": "Dandi March",
    "दांडी": "Dandi",
}

_ask_long_before_glossary = ask_long
def ask_long(*a, **k):
    if a and isinstance(a[0], str) and a[0].startswith("Translate this question to English"):
        head, sep, body = a[0].partition("\n\n")
        for key in sorted(GLOSSARY, key=len, reverse=True):
            body = body.replace(key, GLOSSARY[key])
        a = (head + sep + body,) + a[1:]
    return _ask_long_before_glossary(*a, **k)
# END


# ======== v23 (14 Oct 2026): cleaner "current" answers, page choice with first sentences, refusal in the user's language ========
def current_scan(title, eng, max_chunks=16):
    parts = []
    for c in make_chunks(title)[:max_chunks]:
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c +
                     "\n\nIf this TEXT states an answer, reply with the answer and the year or date it applies to, "
                     "in one line. If not, reply NONE.", 0, n=60).strip()
        if r and not r.upper().startswith("NONE"):
            parts.append(r)
    if not parts:
        return None
    best = ask_long("Question: " + eng + "\n\nCANDIDATE ANSWERS:\n" + "\n".join(parts) +
                    "\n\nChoose the most recent answer. Reply with ONLY that answer and the date or year it applies "
                    "to, in one short line. No explanation, no disclaimer.", 0, n=80).strip().splitlines()[0]
    if re.search(r"knowledge cutoff|real-time|I do not|I cannot", best, re.I):
        best = parts[0]
    return best

def _intros(titles):
    try:
        d = _rq.get(API, params={"action": "query", "format": "json", "prop": "extracts", "exintro": 1,
                                 "explaintext": 1, "exchars": 110, "exlimit": "max", "titles": "|".join(titles)},
                    headers=_UA, timeout=15).json()
        return {p["title"]: " ".join(p.get("extract", "").split()) for p in d["query"]["pages"].values()}
    except Exception:
        return {}

def pick_from_search(eng, seed, exclude=(), extra_query=None, size=6):
    titles = []
    for qq in [eng] + ([extra_query] if extra_query else []):
        for t in search_titles(qq, 8):
            if t not in exclude and t not in titles:
                titles.append(t)
    titles = titles[:size]
    if not titles:
        return None
    intro = _intros(titles)
    menu = "\n".join(f"{i + 1}. {t} - {intro.get(t, '')[:110]}" for i, t in enumerate(titles))
    pick = ask_long("Question: " + eng + "\n\nWikipedia pages:\n" + menu +
                    "\n\nPrefer the main article about the topic, not a book, film or episode unless the question "
                    "asks for one. Reply with only the number of the page most likely to contain the answer.",
                    seed, n=5).strip()
    m = re.search(r"\d", pick)
    i = int(m.group()) - 1 if m else 0
    return titles[i if 0 <= i < len(titles) else 0]

_agent_before_lang = agent
def agent(question, seed=0, max_tries=3):
    out = _agent_before_lang(question, seed, max_tries)
    lang = detect_lang(question)
    if lang != "English" and "(No source checked" in out:
        return NOT_FOUND[lang]
    return out
# END


# ======== v24 (14 Oct 2026): page choice by rules - key-word search + drop film/book/album pages unless asked ========
_STOP = set("list name the of what who is are was were a an in on at which give tell me about all and to for from how many much does do did when where why".split())
_NUM = set("one two three four five six seven eight nine ten".split())
_KIND = re.compile(r"\b(film|novel|television series|TV series|album|video game|song|episode)\b", re.I)
_ASKS = ("film", "movie", "novel", "book", "series", "episode", "album", "song", "game", "direct", "actor",
         "actress", "cast", "released", "plot", "sing", "sung", "author", "wrote", "written", "music", "singer")

def _trim(q):
    return " ".join(w for w in re.findall(r"[\w']+", q)
                    if w.lower() not in _STOP and w.lower() not in _NUM and not w.isdigit())

def pick_from_search(eng, seed, exclude=(), extra_query=None, size=6):
    titles = []
    groups = [search_titles(eng, 8)[:4], search_titles(_trim(eng), 8)[:4] if _trim(eng) else []]
    if extra_query:
        groups.append(search_titles(extra_query, 8)[:2])
    for g in groups:
        for t in g:
            if t not in exclude and t not in titles:
                titles.append(t)
    titles = titles[:max(size, 8)]
    if not titles:
        return None
    intro = _intros(titles)
    if not any(a in eng.lower() for a in _ASKS):
        good = [t for t in titles if not _KIND.search(intro.get(t, ""))]
        titles = good or titles
    menu = "\n".join(f"{i + 1}. {t} - {intro.get(t, '')[:110]}" for i, t in enumerate(titles))
    pick = ask_long("Question: " + eng + "\n\nWikipedia pages:\n" + menu +
                    "\n\nPrefer the main article about the topic. Reply with only the number of the page "
                    "most likely to contain the answer.", seed, n=5).strip()
    m = re.search(r"\d", pick)
    i = int(m.group()) - 1 if m else 0
    return titles[i if 0 <= i < len(titles) else 0]
# END


# ======== v25 (14 Oct 2026): verify the answer line is specific; note other meanings for bare topics ========
_VERIFY = ("Question: {q}\n\nProposed answer: {a}\n\nDoes the proposed answer give the specific information the "
           "question asks for (for example the exact place, date, name or number)? Reply YES or NO.")

def chunk_scan(title, eng, max_chars=4500, max_chunks=12):
    blocks = [b for b in page_blocks2(title) if not b.startswith("^")]
    chunks, cur = [], ""
    for b in blocks:
        if len(cur) + len(b) > max_chars:
            chunks.append(cur)
            cur = ""
        cur += b[:300] + "\n"
    if cur:
        chunks.append(cur)
    first_yes, first_i = None, 0
    for i, c in enumerate(chunks[:max_chunks]):
        if first_yes and i > first_i + 4:
            break
        r = ask_long("Question: " + eng + "\n\nTEXT:\n" + c +
                     "\n\nDoes this TEXT directly state the answer? If yes, reply 'YES: ' followed by the answer. "
                     "If not, reply 'NO'.", 0, n=40).strip()
        if r.upper().startswith("YES"):
            v = ask_long(_VERIFY.format(q=eng, a=r[4:].strip(" :")), 0, n=5).strip()
            if v.upper().startswith("YES"):
                return r
            if not first_yes:
                first_yes, first_i = r, i
    return None
chunk_scan = _wrap_scan(chunk_scan)

_agent_before_note = agent
def agent(question, seed=0, max_tries=3):
    out = _agent_before_note(question, seed, max_tries)
    if out.startswith(tuple(ASK.values())):
        return out
    lang = detect_lang(question)
    base = question if lang == "English" else ask_long(
        "Translate this question to English. Output only the translation.\n\n" + question, seed, n=40).strip()
    topic = re.sub(r"^(please\s+)?(tell me about|what is|what are|who is|who was|explain|about|define)\s+", "",
                   base.strip().rstrip("?.! "), flags=re.I)
    topic = re.sub(r"^(a|an|the)\s+", "", topic, flags=re.I).strip()
    known = any(_norm(f) in _norm(question) or _norm(f) in _norm(topic) for fs in NAME_FORMS.values() for f in fs)
    if topic and len(topic.split()) <= 2 and not known:
        try:
            opts, real = disambig_options(topic.title() + " (disambiguation)")
        except Exception:
            opts = None
        if opts:
            return out + "\n(" + ASK[lang] + " " + "; ".join(opts[:6]) + ")"
    return out
# END


# ======== v26 (14 Oct 2026): 60 s time budget per question; no answer-check for bare topics ========
class _OverBudget(BaseException):
    pass
_DL = {"t": 0.0}
BUDGET_SECONDS = 60

def _bare_topic(q):
    t = re.sub(r"^(please\s+)?(tell me about|what is|what are|who is|who was|explain|about|define|describe)\s+", "",
               q.strip().rstrip("?.! "), flags=re.I)
    t = re.sub(r"^(a|an|the)\s+", "", t, flags=re.I)
    return len(t.split()) <= 2

_ask_before_budget = ask_long
def ask_long(*a, **k):
    if _DL["t"] and time.time() > _DL["t"]:
        raise _OverBudget()
    return _ask_before_budget(*a, **k)

_agent_before_budget = agent
def agent(question, seed=0, max_tries=3):
    low = question.lower()
    if any(w in low for w in CODE_WORDS) or _FILE_RE.search(question):
        return _agent_before_budget(question, seed, max_tries)
    _DL["t"] = time.time() + BUDGET_SECONDS
    try:
        return _agent_before_budget(question, seed, max_tries)
    except _OverBudget:
        return (NOT_FOUND[detect_lang(question)] + f"\n(Stopped after {BUDGET_SECONDS} seconds so you do not "
                "wait too long. Try a more specific question.)")
    finally:
        _DL["t"] = 0.0
# END


# ======== v26b: skip the answer-check for bare topics ========
_ask_before_bare = ask_long
def ask_long(*a, **k):
    p = a[0] if a else ""
    if p.startswith("Question:") and "\n\nProposed answer:" in p:
        q = p.split("\n\nProposed answer:")[0][len("Question:"):].strip()
        if _bare_topic(q):
            return "YES"
    return _ask_before_bare(*a, **k)
# END


# ======== v26c: do not scan the same page twice; budget 90 s ========
BUDGET_SECONDS = 90
_cs_prev = chunk_scan
_cs_cache = {}
def chunk_scan(title, eng, *a, **k):
    key = (title, eng)
    if key not in _cs_cache:
        _cs_cache[key] = _cs_prev(title, eng, *a, **k)
    return _cs_cache[key]
# END

# ======== v26d (15 Oct 2026): cache a page scan by the page's own text, not its title ========
import hashlib
_cs_prev2 = chunk_scan
_cs_cache2 = {}
def _page_key(title, eng):
    try:
        ch = make_chunks(title)
        return (eng, hashlib.md5((ch[0] if ch else title).encode("utf-8")).hexdigest())
    except Exception:
        return (title, eng)
def chunk_scan(title, eng, *a, **k):
    key = _page_key(title, eng)
    if key in _cs_cache2:
        r = _cs_cache2[key]
        if r:
            LAST["title"] = title
        return r
    r = _cs_prev2(title, eng, *a, **k)
    _cs_cache2[key] = r
    return r
# END

# ======== v26e (15 Oct 2026): remember Wikipedia answers 10 min; shorter waits; print the reason ========
_wg_cache = {}
def wiki_get(params, tries=3):
    key = json.dumps(params, sort_keys=True)
    hit = _wg_cache.get(key)
    if hit and time.time() - hit[0] < 600:
        return hit[1]
    last = None
    for i in range(tries):
        try:
            r = requests.get(API, params=params, headers=H, timeout=12)
        except requests.RequestException as e:
            last = "error " + type(e).__name__
            print(f"[wiki] {last}, retry {i + 1}", flush=True)
            time.sleep(1)
            continue
        if r.status_code == 200 and "json" in r.headers.get("content-type", ""):
            data = r.json()
            _wg_cache[key] = (time.time(), data)
            return data
        last = r.status_code
        print(f"[wiki] status {r.status_code}, Retry-After={r.headers.get('Retry-After')}, retry {i + 1}", flush=True)
        try:
            wait = float(r.headers.get("Retry-After", 2 * (i + 1)))
        except ValueError:
            wait = 2 * (i + 1)
        time.sleep(min(wait, 5))
    raise RuntimeError("Wikipedia kept refusing, last status " + str(last))
# END

_wg_last = [0.0]
def wiki_get(params, tries=4):
    key = json.dumps(params, sort_keys=True)
    hit = _wg_cache.get(key)
    if hit and time.time() - hit[0] < 600:
        return hit[1]
    last = None
    for i in range(tries):
        gap = 0.7 - (time.time() - _wg_last[0])
        if gap > 0: time.sleep(gap)
        _wg_last[0] = time.time()
        try:
            r = requests.get(API, params=params, headers=H, timeout=12)
        except requests.RequestException as e:
            last = "error " + type(e).__name__
            print(f"[wiki] {last}, retry {i + 1}", flush=True)
            time.sleep(1); continue
        if r.status_code == 200 and "json" in r.headers.get("content-type", ""):
            data = r.json(); _wg_cache[key] = (time.time(), data); return data
        last = r.status_code
        print(f"[wiki] status {r.status_code}, Retry-After={r.headers.get('Retry-After')}, retry {i + 1}", flush=True)
        try: wait = float(r.headers.get("Retry-After", 3 * (i + 1)))
        except ValueError: wait = 3 * (i + 1)
        time.sleep(min(wait, 30))
    raise RuntimeError("Wikipedia kept refusing, last status " + str(last))

_agent_prev_f = agent
def agent(question, *a, **k):
    try:
        return _agent_prev_f(question, *a, **k)
    except RuntimeError as e:
        if "Wikipedia kept refusing" in str(e):
            return "Wikipedia is limiting our requests right now (too many requests). Please wait a minute and ask again."
        raise
# END

import json, re, os, time, datetime
KB_FILE = "/kaggle/working/kb.json"
LOG_FILE = "/kaggle/working/qa_log.jsonl"
_STOP = set("a an the is are was were be been am do does did of in on at to for from by with and or me my tell about please can you could would should it its this that these those".split())
_QW = {"who", "what", "where", "when", "why", "how", "which", "whom", "whose"}
_CURRENT = set("current currently now today latest chief minister president price ceo prime champion score weather richest".split())
TTL_STABLE_DAYS = 30
TTL_CURRENT_DAYS = 1

def kb_key(eng):
    words = re.findall(r"[a-z0-9]+", eng.lower())
    keep = sorted(set(w for w in words if w in _QW or w not in _STOP))
    return "|".join(keep)

def kb_is_current(eng):
    return any(w in _CURRENT for w in re.findall(r"[a-z]+", eng.lower()))

def kb_load():
    try:
        with open(KB_FILE, encoding="utf-8") as f: return json.load(f)
    except Exception:
        return {}

def kb_save(db):
    tmp = KB_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f: json.dump(db, f, ensure_ascii=False, indent=1)
    os.replace(tmp, KB_FILE)

def kb_get(db, eng, lang, now=None):
    now = now or time.time()
    e = db.get(kb_key(eng) + "#" + lang)
    if not e: return None
    ttl = TTL_CURRENT_DAYS if e.get("current") else TTL_STABLE_DAYS
    if now - e["checked_ts"] > ttl * 86400: return None
    return e

def kb_put(db, eng, lang, answer, source, title="", chunk="", edited="", passed=False, now=None):
    if not passed: return False
    now = now or time.time()
    db[kb_key(eng) + "#" + lang] = {
        "q": eng, "lang": lang, "answer": answer, "source": source, "title": title,
        "chunk": chunk[:4500], "edited": edited, "current": kb_is_current(eng),
        "checked_ts": now, "checked": datetime.datetime.fromtimestamp(now, datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")}
    return True

def log_add(question, eng, lang, answer, source, passed, seconds, from_cache=False):
    row = {"ts": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M:%S"), "question": question, "eng": eng,
           "key": kb_key(eng), "lang": lang, "answer": answer, "source": source,
           "passed": bool(passed), "seconds": round(seconds, 1), "from_cache": from_cache}
    with open(LOG_FILE, "a", encoding="utf-8") as f: f.write(json.dumps(row, ensure_ascii=False) + "\n")
# END

KB = kb_load()
_agent_prev_kb = agent
_NOTE = {"English": "(Saved answer, checked {d})",
         "Hindi": "(सहेजा हुआ उत्तर, जाँचा गया: {d})",
         "Telugu": "(సేవ్ చేసిన సమాధానం, తనిఖీ: {d})"}

def _kb_skip(question):
    try:
        if _FILE_RE.search(question): return True
    except Exception: pass
    try:
        ql = question.lower()
        if any(w in ql for w in CODE_WORDS): return True
    except Exception: pass
    return False

def _kb_good(text):
    if "(Source:" not in text: return False
    if any(v in text for v in NOT_FOUND.values()): return False
    if "Stopped after" in text or "limiting our requests" in text or "last edited unknown" in text: return False
    return True

def agent(question, *a, **k):
    t0 = time.time()
    if _kb_skip(question):
        return _agent_prev_kb(question, *a, **k)
    lang = detect_lang(question)
    if lang == "English":
        eng = question
    else:
        eng = str(ask_long("Translate this question to English. Output only the translation.\n\n" + question, 0, 80)).strip() or question
    e = kb_get(KB, eng, lang)
    if e:
        print("[kb] saved answer used, key:", kb_key(eng), flush=True)
        ans = e["answer"] + "\n" + _NOTE[lang].format(d=e["checked"])
        log_add(question, eng, lang, ans, e.get("source", ""), True, time.time() - t0, True)
        return ans
    ans = str(_agent_prev_kb(question, *a, **k))
    good = _kb_good(ans)
    m = re.search(r"\(Source:[^\n]*\)", ans)
    if good:
        kb_put(KB, eng, lang, ans, m.group(0) if m else "", passed=True)
        kb_save(KB)
        print("[kb] saved new answer, key:", kb_key(eng), flush=True)
    else:
        print("[kb] not saved (failed checks)", flush=True)
    log_add(question, eng, lang, ans, m.group(0) if m else "", good, time.time() - t0, False)
    return ans
# END

def page_edited(title):
    try:
        d = wiki_get({"action": "query", "titles": title, "prop": "revisions",
                      "rvprop": "timestamp", "redirects": 1, "format": "json"}, tries=2)
    except Exception as e:
        print("[wiki] edit date failed:", str(e)[:80], flush=True)
        return "unknown"
    for pg in d.get("query", {}).get("pages", {}).values():
        revs = pg.get("revisions")
        if revs:
            return revs[0]["timestamp"]
    return "unknown"
# END

_WEAK = ("Only one source found", "No source checked", "model's memory", "not cross-checked")
def _kb_good(text):
    if "(Source:" not in text: return False
    if any(v in text for v in NOT_FOUND.values()): return False
    if "Stopped after" in text or "limiting our requests" in text or "last edited unknown" in text: return False
    if any(w in text for w in _WEAK): return False
    return True
# END


# ======== v27d_kb.py ========
import json, os
_BLOCK_FILE = "/kaggle/working/kb_block.json"
_block_init = [
 "closest|earth|star|what#English", "closest|earth|star|what#Hindi", "closest|earth|star|what#Telugu",
 "first|general|governor|independent|india|who#English", "first|general|governor|independent|india|who#Hindi",
 "economics|who|wrote#Hindi",
 "composed|film|millionaire|music|slumdog|who#Telugu",
 "2016|country|hosted|olympics|summer|which#Telugu",
]
json.dump(_block_init, open(_BLOCK_FILE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

def block_load():
    try:
        return set(json.load(open(_BLOCK_FILE, encoding="utf-8")))
    except Exception:
        return set()
BLOCK = block_load()

_NOTE2 = {"English": "(Saved answer, not yet reviewed by a person, checked {d})",
          "Hindi": "(सहेजा हुआ उत्तर, अभी किसी व्यक्ति ने नहीं जाँचा; जाँचा गया: {d})",
          "Telugu": "(సేవ్ చేసిన సమాధానం, ఇంకా ఏ వ్యక్తీ తనిఖీ చేయలేదు; తనిఖీ: {d})"}
_WARN = {"English": "(Warning: this answer was found wrong in our tests. Please check it yourself.)",
         "Hindi": "(चेतावनी: हमारे परीक्षण में यह उत्तर गलत पाया गया था। कृपया स्वयं जाँच लें।)",
         "Telugu": "(హెచ్చరిక: మా పరీక్షలో ఈ సమాధానం తప్పుగా తేలింది. దయచేసి మీరే సరిచూసుకోండి.)"}

_kb_put_prev = kb_put
def kb_put(db, eng, lang, answer, source, title="", chunk="", edited="", passed=False, now=None):
    ok = _kb_put_prev(db, eng, lang, answer, source, title=title, chunk=chunk, edited=edited, passed=passed, now=now)
    if ok:
        db[kb_key(eng) + "#" + lang]["reviewed"] = False
    return ok

def agent(question, *a, **k):
    t0 = time.time()
    if _kb_skip(question):
        return _agent_prev_kb(question, *a, **k)
    lang = detect_lang(question)
    if lang == "English":
        eng = question
    else:
        eng = str(ask_long("Translate this question to English. Output only the translation.\n\n" + question, 0, 80)).strip() or question
    full = kb_key(eng) + "#" + lang
    if full in BLOCK:
        ans = str(_agent_prev_kb(question, *a, **k))
        print("[kb] blocked key (known wrong), not saved:", kb_key(eng), flush=True)
        ans = ans + "\n" + _WARN[lang]
        log_add(question, eng, lang, ans, "", False, time.time() - t0, False)
        return ans
    e = kb_get(KB, eng, lang)
    if e:
        print("[kb] saved answer used, key:", kb_key(eng), flush=True)
        note = (_NOTE if e.get("reviewed", True) else _NOTE2)[lang].format(d=e["checked"])
        ans = e["answer"] + "\n" + note
        log_add(question, eng, lang, ans, e.get("source", ""), True, time.time() - t0, True)
        return ans
    ans = str(_agent_prev_kb(question, *a, **k))
    good = _kb_good(ans)
    m = re.search(r"\(Source:[^\n]*\)", ans)
    if good:
        kb_put(KB, eng, lang, ans, m.group(0) if m else "", passed=True)
        kb_save(KB)
        print("[kb] saved new answer, key:", kb_key(eng), flush=True)
    else:
        print("[kb] not saved (failed checks)", flush=True)
    log_add(question, eng, lang, ans, m.group(0) if m else "", good, time.time() - t0, False)
    return ans


# ======== v27d_clean.py ========

import re

_SUB = str.maketrans("0123456789+-", "₀₁₂₃₄₅₆₇₈₉₊₋")
_SUP = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")

def _fix_math(m):
    s = m.group(1)
    if "\\" not in s and "^" not in s and "_" not in s:
        return m.group(0)
    s = re.sub(r"\\(?:text|mathrm|mathbf|textbf)\{([^}]*)\}", r"\1", s)
    s = s.replace("\\circ", "°").replace("\\Omega", "Ω").replace("\\times", "×").replace("\\mu", "µ")
    s = re.sub(r"_\{?([0-9+\-]+)\}?", lambda x: x.group(1).translate(_SUB), s)
    s = re.sub(r"\^\{?([0-9+\-]+)\}?", lambda x: x.group(1).translate(_SUP), s)
    s = s.replace("^°", "°").replace("° ", "°")
    s = s.replace("{", "").replace("}", "").replace("\\", "")
    return s.strip()

def clean_answer(text):
    text = re.sub(r"\$([^$\n]{1,80})\$", _fix_math, str(text))
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    return text

if not globals().get("_clean_wrapped"):
    _agent_before_clean = agent
    def agent(question, *a, **k):
        return clean_answer(str(_agent_before_clean(question, *a, **k)))
    _clean_wrapped = True


# ======== v27d_route.py ========

import re
if not globals().get("_route_fixed"):
    _CW_ALL = CODE_WORDS
    def _code_words_for(question):
        low = question.lower()
        if "python programming" in low and not re.search(r"python program(?!ming)", low):
            return tuple(w for w in _CW_ALL if w != "python program")
        return _CW_ALL
    _agent_before_route_fix = agent
    def agent(question, *a, **k):
        global CODE_WORDS
        CODE_WORDS = _code_words_for(question)
        try:
            return _agent_before_route_fix(question, *a, **k)
        finally:
            CODE_WORDS = _CW_ALL
    _route_fixed = True


# ======== v27d_reviewed.py ========

if not globals().get("_pin_ready"):
    _kb_get_prev_pin = kb_get
    def kb_get(db, eng, lang, now=None):
        e = db.get(kb_key(eng) + "#" + lang)
        if e and e.get("pinned"):
            return e
        return _kb_get_prev_pin(db, eng, lang, now)
    _pin_ready = True

def kb_add_reviewed(eng, lang, answer, source, reviewed):
    now = time.time()
    k = kb_key(eng) + "#" + lang
    KB[k] = {"q": eng, "lang": lang, "answer": answer + "\n" + source, "source": source, "title": "", "chunk": "",
             "edited": "", "current": False, "pinned": True, "reviewed": bool(reviewed), "checked_ts": now,
             "checked": datetime.datetime.fromtimestamp(now, datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")}
    BLOCK.discard(k)
    kb_save(KB)
    json.dump(sorted(BLOCK), open(_BLOCK_FILE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return k


def load_reviewed(path):
    n = 0
    for r in json.load(open(path, encoding="utf-8")):
        kb_add_reviewed(r["q"], r["lang"], r["answer"], r["source"], r["reviewed"])
        n += 1
    return n



# ======== long answers EN ========

import re, time
LA_SKIP = {"references", "see also", "external links", "notes", "further reading", "bibliography",
           "sources", "citations", "footnotes", "gallery", "works cited"}
_LA_IPA = "ˈˌːəɑɪʊʃʒθðŋɾʈɖʰɦɛɔ̪͡"
_la_pron = ("It ", "He ", "She ", "They ", "This ", "These ", "His ", "Her ", "Its ", "Their ", "There ", "That ", "Those ")

def _la_sections(title):
    j = wiki_get({"action": "parse", "format": "json", "formatversion": 2, "prop": "text", "page": title, "redirects": 1})
    soup = BeautifulSoup(j["parse"]["text"], "html.parser")
    for t in soup(["style", "script", "sup"]):
        t.decompose()
    for t in soup.select(".mw-editsection"):
        t.decompose()
    sections, parent = [["Lead", [], "lead", "Lead"]], "Lead"
    for el in soup.find_all(["h2", "h3", "p"]):
        if el.name in ("h2", "h3"):
            name = re.sub(r"\s+", " ", el.get_text(" ", strip=True)).strip()
            if el.name == "h2":
                parent = name
            sections.append([name, [], el.name, parent])
        else:
            s = re.sub(r"\s+", " ", el.get_text(" ", strip=True))
            if len(s) >= 120:
                sections[-1][1].append(s)
    return sections

def _la_sentences(par):
    raw = re.split(r"(?<=[.!?])\s+(?=[A-Z])", par.strip())
    merged = []
    for s in raw:
        if merged and re.search(r"(?:\b[A-Za-z]|\bDr|\bMr|\bMrs|\bSt|\bNo|\bca)\.$", merged[-1]):
            merged[-1] += " " + s
        else:
            merged.append(s)
    return [s.strip() for s in merged]

def _la_sent_ok(s, allow_pron=False):
    if not (40 <= len(s) <= 420) or not s[0].isupper():
        return False
    if not re.search(r'[.!?]["\')\]]?$', s):
        return False
    if s.count("(") != s.count(")") or s.count("[") != s.count("]") or s.count('"') % 2 == 1:
        return False
    if "[...]" in s or "..." in s:
        return False
    return allow_pron or not s.startswith(_la_pron)

def _la_is_prose(par):
    if len(par) < 120 or not par[0].isupper() or "[...]" in par or par.count(":") > 4:
        return False
    commas = par.count(",")
    return not (commas >= 6 and len(par) / (commas + 1) < 20)

def _la_clean(s):
    s = s.replace("\u2060", "").replace("\u200b", "")
    s = re.sub(r"\[[^\]]*[" + _LA_IPA + r"][^\]]*\]", "", s)
    s = re.sub(r"\b[A-Z][a-z]+:\s*(?=[;)])", "", s)
    s = re.sub(r"\(\s*/[^;)]*?/[^;)]*;\s*", "(", s)
    s = re.sub(r"\(\s*;\s*", "(", s)
    s = re.sub(r";\s*;", ";", s)
    s = re.sub(r"\s+([,.;:)\]])", r"\1", s)
    s = re.sub(r"([(\[])\s+", r"\1", s)
    s = re.sub(r"(\w)- (\w)", r"\1-\2", s)
    s = re.sub(r"\s'\s+([^']+?)\s+'", r" '\1'", s)
    s = re.sub(r"\(\s*\)", "", s)
    return re.sub(r"\s{2,}", " ", s).strip()

def long_answer(topic, words=250, seed=0):
    t0 = time.time()
    title = pick_from_search(topic, seed)
    if not title:
        return {"ok": False, "text": NOT_FOUND["English"]}
    secs = _la_sections(title)
    lead = [p for p in secs[0][1] if _la_is_prose(p)][:2]
    fam, order, total = {}, [], {}
    for s in secs[1:]:
        if s[0].lower() in LA_SKIP or s[3].lower() in LA_SKIP:
            continue
        prose = [p for p in s[1] if _la_is_prose(p)]
        if not prose:
            continue
        total[s[3]] = total.get(s[3], 0) + sum(len(p) for p in prose)
        if s[3] not in fam:
            fam[s[3]] = (s[0], prose[0], len(order))
            order.append(s[3])
    ranked = sorted(order, key=lambda p: (-total[p], fam[p][2]))
    def wc(keep):
        return len(_la_clean(" ".join(keep)).split())
    def pick_sents(par, allow_pron):
        sents = [x for x in _la_sentences(par) if _la_sent_ok(x, allow_pron)]
        if len(sents) <= 2:
            return sents
        numbered = "\n".join(f"{i + 1}. {x}" for i, x in enumerate(sents[:8]))
        r = ask_long("Question: Give a short overview of the topic '" + topic + "'.\n\nHere are numbered sentences from one section. "
                     "Reply ONLY with the numbers of the 1 or 2 most informative sentences that make sense on their own. "
                     "Example reply: 1, 3\n\n" + numbered, seed, n=12).strip()
        nums = [int(x) for x in re.findall(r"\d+", r)][:2]
        return [sents[n - 1] for n in nums if 1 <= n <= len(sents[:8])] or sents[:1]
    used, picked = 0, []
    for p in lead:
        keep = pick_sents(p, True)
        if keep:
            picked.append(("Lead", keep, -1)); used += wc(keep)
    for parent in ranked:
        if used >= words * 0.9:
            break
        keep = pick_sents(fam[parent][1], False)
        if not keep:
            continue
        w = wc(keep)
        if used + w > words * 1.15:
            keep = keep[:1]; w = wc(keep)
            if used + w > words * 1.15:
                continue
        picked.append((parent, keep, fam[parent][2])); used += w
    picked.sort(key=lambda x: x[2])
    page_text = " ".join(" ".join(s[1]) for s in secs)
    exact = all(x in page_text for _, keep, _ in picked for x in keep)
    parts = [(("" if n == "Lead" else n + "\n") + _la_clean(" ".join(keep))) for n, keep, _ in picked]
    try:
        ed = page_edited(title)[:10]
    except Exception:
        ed = "unknown"
    text = "\n\n".join(parts) + f"\n\n(Source: Wikipedia, {title}, last edited {ed}. Text copied from Wikipedia under CC BY-SA; pronunciation guides and stray spaces tidied.)"
    return {"ok": True, "text": text, "title": title, "exact": exact, "seconds": round(time.time() - t0, 1), "n": len(picked),
            "chosen": [(n, total.get(n, 0)) for n, _, _ in picked if n != "Lead"],
            "left_out": [(n, total[n]) for n in ranked if n not in [x[0] for x in picked]][:6]}

# ---- v3 ----

_la_conn = ("One result", "One such", "Another ", "However", "Thus", "Therefore", "Hence", "Moreover", "Furthermore",
            "Meanwhile", "In addition", "Also,", "As a result", "For this reason")

def _la_sent_ok(s, allow_pron=False):
    if not (40 <= len(s) <= 420) or not s[0].isupper():
        return False
    if not re.search(r'[.!?]["\')\]]?$', s):
        return False
    if s.count("(") != s.count(")") or s.count("[") != s.count("]") or s.count('"') % 2 == 1:
        return False
    if "[...]" in s or "..." in s:
        return False
    if s.startswith(_la_conn):
        return False
    return allow_pron or not s.startswith(_la_pron)

def _la_pick(topic, label, sents, n_max, seed):
    sents = sents[:10]
    if len(sents) <= n_max:
        return sents
    numbered = "\n".join(f"{i + 1}. {x}" for i, x in enumerate(sents))
    r = ask_long(f"Question: Give a short overview of the topic '{topic}'.\n\nHere are numbered sentences from the part '{label}'. "
                 f"Reply ONLY with the numbers of the {n_max} most informative sentences that explain the topic and make sense on their own. "
                 "Example reply: 1, 3\n\n" + numbered, seed, n=12).strip()
    nums = []
    for x in re.findall(r"\d+", r):
        v = int(x)
        if 1 <= v <= len(sents) and v not in nums:
            nums.append(v)
    nums = sorted(nums[:n_max])
    return [sents[v - 1] for v in nums] or sents[:1]

def long_answer(topic, words=250, seed=0):
    t0 = time.time()
    title = pick_from_search(topic, seed)
    if not title:
        return {"ok": False, "text": NOT_FOUND["English"]}
    secs = _la_sections(title)
    fam, order, total = {}, [], {}
    for s in secs[1:]:
        if s[0].lower() in LA_SKIP or s[3].lower() in LA_SKIP:
            continue
        prose = [p for p in s[1] if _la_is_prose(p)]
        if not prose:
            continue
        total[s[3]] = total.get(s[3], 0) + sum(len(p) for p in prose)
        if s[3] not in fam:
            fam[s[3]] = [s[0], [], len(order)]
            order.append(s[3])
        if len(fam[s[3]][1]) < 2:
            fam[s[3]][1].extend(prose[:2 - len(fam[s[3]][1])])
    names = order[:30]
    k_sec = max(2, min(7, round(words / 45) - 2))
    numbered = "\n".join(f"{i + 1}. {n[:50]}" for i, n in enumerate(names))
    r = ask_long(f"Topic: {topic}. Here are the section headings of its encyclopedia article.\n\n{numbered}\n\n"
                 f"Reply ONLY with the numbers of the {k_sec} headings most important for a general overview of the topic "
                 "(prefer the core subject matter; avoid minor, trivia, research-history or list sections). Example reply: 2, 4, 5, 7",
                 seed, n=24).strip()
    chosen_idx = []
    for x in re.findall(r"\d+", r):
        v = int(x)
        if 1 <= v <= len(names) and v not in chosen_idx:
            chosen_idx.append(v)
    chosen = [names[v - 1] for v in chosen_idx[:k_sec]]
    for n in sorted(names, key=lambda p: -total[p]):
        if len(chosen) >= k_sec:
            break
        if n not in chosen:
            chosen.append(n)
    chosen.sort(key=lambda n: fam[n][2])
    def wc(keep):
        return len(_la_clean(" ".join(keep)).split())
    lead_pars = [p for p in secs[0][1] if _la_is_prose(p)][:2]
    lead_sents = []
    for i, p in enumerate(lead_pars):
        for j, x in enumerate(_la_sentences(p)):
            if _la_sent_ok(x, allow_pron=(i == 0 and j == 0)):
                lead_sents.append(x)
    lead_keep = _la_pick(topic, "Lead", lead_sents, 3, seed)
    picked, used = [("Lead", lead_keep, -1)], wc(lead_keep)
    while used > words * 0.45 and len(picked[0][1]) > 1:
        picked[0] = ("Lead", picked[0][1][:-1], -1)
        used = wc(picked[0][1])
    per = max(1, round((words - used) / max(1, len(chosen)) / 28))
    for n in chosen:
        if used >= words * 1.05:
            break
        sents = [x for par in fam[n][1] for x in _la_sentences(par) if _la_sent_ok(x, False)]
        if not sents:
            continue
        keep = _la_pick(topic, fam[n][0], sents, min(2, per), seed)
        w = wc(keep)
        if used + w > words * 1.2:
            keep = keep[:1]; w = wc(keep)
            if used + w > words * 1.2:
                continue
        picked.append((n, keep, fam[n][2])); used += w
    picked.sort(key=lambda x: x[2])
    page_text = " ".join(" ".join(s[1]) for s in secs)
    exact = all(x in page_text for _, keep, _ in picked for x in keep)
    parts = [(("" if n == "Lead" else n + "\n") + _la_clean(" ".join(keep))) for n, keep, _ in picked]
    try:
        ed = page_edited(title)[:10]
    except Exception:
        ed = "unknown"
    text = "\n\n".join(parts) + f"\n\n(Source: Wikipedia, {title}, last edited {ed}. Text copied from Wikipedia under CC BY-SA; pronunciation guides and stray spaces tidied.)"
    body_words = len(re.sub(r"\(Source:.*$", "", text, flags=re.S).split())
    return {"ok": True, "text": text, "title": title, "exact": exact, "seconds": round(time.time() - t0, 1), "n": len(picked),
            "body_words": body_words, "model_chose": [names[v - 1] for v in chosen_idx[:k_sec]],
            "used_headings": [n for n, _, _ in picked if n != "Lead"]}

# ---- v4: request detection, agent_long() ----

_LA_NUM = re.compile(r"\b(\d{2,4})\s*[- ]?\s*words?\b", re.I)
_LA_LONG = re.compile(r"\b(in detail|detailed|in depth|in-depth|elaborate|comprehensive|long answer|essay|full explanation|explain fully|overview of|tell me everything|at length)\b", re.I)
_LA_STRIP = re.compile(r"\b(please|can you|could you|give me|write|explain|describe|tell me|about|in detail|detailed|in depth|in-depth|an? essay on|essay on|an? overview of|overview of|long answer|at length|elaborate on|comprehensive|everything|in\s+\d{2,4}\s*[- ]?\s*words?|\d{2,4}\s*[- ]?\s*words?|the|a|an|on)\b", re.I)

def la_request(q):
    if detect_lang(q) != "English":
        return None
    m = _LA_NUM.search(q)
    if not m and not _LA_LONG.search(q):
        return None
    words = max(100, min(int(m.group(1)) if m else 250, 600))
    topic = _LA_STRIP.sub(" ", q)
    topic = re.sub(r"[?.!,:;]+", " ", topic)
    topic = re.sub(r"\s+", " ", topic).strip()
    return (words, topic) if len(topic) >= 3 else None

def agent_long(q, *a, **k):
    r = la_request(q)
    if not r:
        return agent(q, *a, **k)
    words, topic = r
    res = long_answer(topic, words)
    return res["text"] if res.get("ok") else agent(q, *a, **k)



# ======== long answers EN v5.5 ========

# ---- v5: length fit (within 10 %), more tidy rules, safer request detection, safe agent_long ----
import re, time

if "_la_clean_v4" not in globals():
    _la_clean_v4 = _la_clean
if "_la_sent_ok_v3" not in globals():
    _la_sent_ok_v3 = _la_sent_ok
if "agent_base" not in globals():
    agent_base = agent

def _la_clean(s):
    s = re.sub(r"\(\s*/[^();]*[" + _LA_IPA + r"][^();]*\)", "", s)
    s = re.sub(r"\(\s*([A-Za-z][\w\- ]*?)\s*/[^();/]*[" + _LA_IPA + r"][^();/]*/\s*\)", r"(\1)", s)
    s = _la_clean_v4(s)
    s = re.sub(r"\s*\(\s*[A-Z]{2,5}:\s*\)", "", s)
    return re.sub(r"\s{2,}", " ", s).strip()

_la_dang_start = ("Later", "Subsequently", "Afterwards", "Afterward", "Eventually", "Soon ", "Then ", "Next,",
                  "Previously", "Earlier", "Initially", "Originally", "Finally", "Consequently", "Nevertheless",
                  "Nonetheless", "Despite this", "Despite that", "In response", "On the other hand",
                  "At the same time", "Because of this", "For example", "For instance", "Similarly", "Likewise",
                  "Conversely", "Instead", "Still,", "Yet ", "But ", "And ", "So ")
_la_dang_re = re.compile(r"^(The|This|That) [\w\- ]{0,30}? (was|were|is|are|had|has) (later|then|subsequently|also|still)\b")
_la_dang_in = (" the former", " the latter", " the above", " as follows", " the aforementioned", " the said ")

def _la_sent_ok(s, allow_pron=False):
    if not _la_sent_ok_v3(s, allow_pron):
        return False
    if s.startswith(_la_dang_start) or _la_dang_re.search(s):
        return False
    low = s.lower()
    return not any(k in low for k in _la_dang_in)

_LA_FIX = {}

def _la_repair(x):
    # one unclosed "(" in the first sentence: cut from the "(" up to the main verb
    if x.count("(") - x.count(")") != 1:
        return None
    i = x.rfind("(")
    m = re.search(r"\s(is|was|are|were)\s", x[i:])
    if not m:
        return None
    y = x[:i].rstrip() + x[i + m.start():]
    return y if _la_sent_ok(_la_clean(y), True) else None

def _la_text(sel, pool, order):
    parts = []
    for n in sorted(sel, key=lambda k: order[k]):
        if not sel[n]:
            continue
        body = _la_clean(" ".join(_LA_FIX.get(pool[n][i], pool[n][i]) for i in sorted(sel[n])))
        parts.append(body if n == "Lead" else n + "\n" + body)
    return "\n\n".join(parts)

def _la_fit(sel, pool, order, words, allowed=None, lo=0.90, hi=1.10, protect=None):
    def size(s):
        return len(_la_text(s, pool, order).split())
    cur = size(sel)
    for _ in range(30):                      # too long: drop the sentence that gets us closest
        if cur <= hi * words:
            break
        best = None
        for n in list(sel):
            for i in list(sel[n]):
                if n == "Lead" and len(sel[n]) <= 1:
                    continue
                if protect is not None and n == protect[0] and i == protect[1]:
                    continue
                s2 = {k: list(v) for k, v in sel.items()}
                s2[n].remove(i)
                if not s2[n]:
                    del s2[n]
                z = size(s2)
                if abs(z - words) < abs(cur - words) and (best is None or abs(z - words) < best[0]):
                    best = (abs(z - words), s2, z)
        if not best:
            break
        sel, cur = best[1], best[2]
    for _ in range(14):                      # too short: add the sentence that gets us closest
        if cur >= lo * words:
            break
        best = None
        for n, sents in pool.items():
            if allowed is not None and n not in allowed:
                continue
            for i in range(len(sents)):
                if n in sel and i in sel[n]:
                    continue
                s2 = {k: list(v) for k, v in sel.items()}
                s2.setdefault(n, []).append(i)
                z = size(s2)
                if z > hi * words:
                    continue
                score = abs(z - words) + (0 if n in sel else 8)
                if abs(z - words) < abs(cur - words) and (best is None or score < best[0]):
                    best = (score, s2, z)
        if not best:
            break
        sel, cur = best[1], best[2]
    return sel, cur

def _la_title_fallback(topic, skip=None):
    try:
        cands = search_titles(topic, 8)
    except Exception as e:
        print("   [title search failed:", type(e).__name__, str(e)[:60] + "]", flush=True)
        return None
    ws = [w for w in re.findall(r"[a-z0-9]+", topic.lower()) if len(w) >= 3]
    if not ws:
        return None
    for t in list(cands)[:5]:
        tl = str(t).lower()
        if skip and tl == str(skip).lower():
            continue
        if all(w in tl for w in ws):
            return t
    return None

def long_answer(topic, words=250, seed=0):
    t0 = time.time()
    title = pick_from_search(topic, seed) or _la_title_fallback(topic)
    if not title:
        return {"ok": False, "text": NOT_FOUND["English"]}
    secs = _la_sections(title)
    if not any(_la_is_prose(p) for p in secs[0][1]):
        alt = _la_title_fallback(topic, skip=title)
        if alt:
            title, secs = alt, _la_sections(alt)
    fam, order, total = {}, [], {}
    for s in secs[1:]:
        if s[0].lower() in LA_SKIP or s[3].lower() in LA_SKIP:
            continue
        prose = [p for p in s[1] if _la_is_prose(p)]
        if not prose:
            continue
        total[s[3]] = total.get(s[3], 0) + sum(len(p) for p in prose)
        if s[3] not in fam:
            fam[s[3]] = [s[0], [], len(order)]
            order.append(s[3])
        if len(fam[s[3]][1]) < 2:
            fam[s[3]][1].extend(prose[:2 - len(fam[s[3]][1])])
    names = order[:30]
    k_sec = max(2, min(7, round(words / 45) - 2))
    numbered = "\n".join(f"{i + 1}. {n[:50]}" for i, n in enumerate(names))
    r = ask_long(f"Topic: {topic}. Here are the section headings of its encyclopedia article.\n\n{numbered}\n\n"
                 f"Reply ONLY with the numbers of the {k_sec} headings most important for a general overview of the topic "
                 "(prefer the core subject matter; avoid minor, trivia, research-history or list sections). Example reply: 2, 4, 5, 7",
                 seed, n=24).strip()
    chosen_idx = []
    for x in re.findall(r"\d+", r):
        v = int(x)
        if 1 <= v <= len(names) and v not in chosen_idx:
            chosen_idx.append(v)
    chosen = [names[v - 1] for v in chosen_idx[:k_sec]]
    for n in sorted(names, key=lambda p: -total[p]):
        if len(chosen) >= k_sec:
            break
        if n not in chosen:
            chosen.append(n)
    chosen.sort(key=lambda n: fam[n][2])
    lead_pars = [p for p in secs[0][1] if _la_is_prose(p)][:2]
    lead_sents = []
    for i, p in enumerate(lead_pars):
        for j, x in enumerate(_la_sentences(p)):
            if _la_sent_ok(x, allow_pron=(i == 0 and j == 0)):
                lead_sents.append(x)
            elif i == 0 and j == 0:
                y = _la_repair(x)
                if y:
                    _LA_FIX[x] = y
                    lead_sents.append(x)
    pool = {"Lead": lead_sents}
    for n in names:
        pool[n] = [x for par in fam[n][1] for x in _la_sentences(par) if _la_sent_ok(x, False)]
    ord_ = {"Lead": -1}
    ord_.update({n: fam[n][2] for n in names})
    lead_keep = _la_pick(topic, "Lead", lead_sents, 3, seed)
    sel = {"Lead": sorted(set([lead_sents.index(x) for x in lead_keep] + [0]))} if lead_sents else {}
    per = max(1, round(words / max(1, len(chosen)) / 40))
    for n in chosen:
        sents = pool[n]
        if not sents:
            continue
        keep = _la_pick(topic, fam[n][0], sents, min(2, per), seed)
        sel[n] = [sents.index(x) for x in keep]
    sel = {k: v for k, v in sel.items() if v}
    if not sel:
        return {"ok": False, "text": NOT_FOUND["English"]}
    sel, cur = _la_fit(sel, pool, ord_, words, allowed=set(["Lead"] + chosen), protect=("Lead", 0))
    if cur < 0.90 * words:
        sel, cur = _la_fit(sel, pool, ord_, words, allowed=None, protect=("Lead", 0))
    page_text = " ".join(" ".join(s[1]) for s in secs)
    exact = all(pool[n][i] in page_text for n in sel for i in sel[n])
    extra = "; a broken bracket of alternative names removed from the first sentence" if any(pool["Lead"][i] in _LA_FIX for i in sel.get("Lead", [])) else ""
    body = _la_text(sel, pool, ord_)
    try:
        ed = page_edited(title)[:10]
    except Exception:
        ed = "unknown"
    text = body + f"\n\n(Source: Wikipedia, {title}, last edited {ed}. Text copied from Wikipedia under CC BY-SA; pronunciation guides and stray spaces tidied{extra}.)"
    return {"ok": True, "text": text, "title": title, "exact": exact, "seconds": round(time.time() - t0, 1),
            "n": len(sel), "body_words": len(body.split()), "target": words,
            "model_chose": [names[v - 1] for v in chosen_idx[:k_sec]],
            "used_headings": [n for n in sorted(sel, key=lambda k: ord_[k]) if n != "Lead"]}

# ---- request detection v5 ----
_LA_LONG = re.compile(r"\b(in detail|in depth|in-depth|elaborate|long answer|full explanation|explain fully|overview of|tell me everything|at length|(?:write|give me|an?) (?:an )?essay on|detailed (?:answer|explanation|overview|description|account))\b", re.I)
_LA_QWORD = re.compile(r"^(how|why|when|where|which|whom|whose|is|are|do|does|did|can|will|would|should)\b", re.I)
_LA_LEAD = re.compile(r"^\s*(what is|what are|what was|what were|who is|who was|who were)\b", re.I)

def la_request(q):
    if detect_lang(q) != "English":
        return None
    m = _LA_NUM.search(q)
    if not m and not _LA_LONG.search(q):
        return None
    words = max(100, min(int(m.group(1)) if m else 250, 600))
    topic = _LA_LEAD.sub(" ", q)
    topic = _LA_STRIP.sub(" ", topic)
    topic = re.sub(r"[?.!,:;]+", " ", topic)
    topic = re.sub(r"\s+", " ", topic).strip()
    if _LA_QWORD.match(topic) or not (3 <= len(topic) <= 60) or len(topic.split()) > 8:
        return None
    return (words, topic)

_NOTE_SHORT = "(A long answer could not be built from Wikipedia for this topic, so this is the short answer instead.)\n\n"

def agent_long(q, *a, **k):
    try:
        r = la_request(q)
    except Exception:
        r = None
    if r:
        words, topic = r
        try:
            res = long_answer(topic, words)
            if res.get("ok") and res.get("exact") and res.get("body_words", 0) >= 60:
                return res["text"]
            print(f"   [long answer not usable: ok={res.get('ok')} exact={res.get('exact')} words={res.get('body_words')}]", flush=True)
        except Exception as e:
            print("   [long answer failed:", type(e).__name__, str(e)[:80] + "]", flush=True)
        return _NOTE_SHORT + str(agent_base(q, *a, **k))
    return agent_base(q, *a, **k)



# ======== Hindi long answers ========
# Hindi long answers, complete experimental version (16 Oct 2026). Needs: agent_v27d, long_answer_experimental (v4), long_answer_v5_5_patch, then this file.
import re, time
HI_API = 'https://hi.wikipedia.org/w/api.php'
HI_SKIP = {'ग्रंथसूची', 'इन्हें भी देखें', 'स्रोत', 'सन्दर्भ सूची', 'टिप्पणियाँ', 'सन्दर्भ', 'संदर्भ', 'बाहरी कड़ी', 'बाहरी कड़ियां', 'टिप्पणी', 'इन्हें भी देखे', 'आगे पढ़ें', 'बाहरी कड़ियाँ', 'नोट', 'ग्रन्थसूची'}
_HI_BACK = ('वह ', 'वे ', 'उन्होंने', 'उसने', 'उन्हें', 'उसे ', 'उसके', 'तब ', 'फिर ', 'इसलिए', 'इसका समाधान', 'इस प्रकार', 'इसके बाद', 'इन ', 'उस ', 'उन ')
_DEV = {2406: 48, 2407: 49, 2408: 50, 2409: 51, 2410: 52, 2411: 53, 2412: 54, 2413: 55, 2414: 56, 2415: 57}
_HI_LONG = re.compile('(विस्तार\\s*से|विस्तृत|विस्तारपूर्वक|निबंध|निबन्ध|लंबा उत्तर|लम्बा उत्तर|पूरी जानकारी|सम्पूर्ण जानकारी|संपूर्ण जानकारी|विस्तार में|\\d+\\s*शब्द(?:ों)?)')
_HI_WORDS = re.compile('(\\d+)\\s*शब्द(?:ों)?')
_HI_NOISE = re.compile('(विस्तार\\s*से|विस्तारपूर्वक|विस्तार में|विस्तृत|निबंध|निबन्ध|लंबा|लम्बा|उत्तर|पूरी|सम्पूर्ण|संपूर्ण|जानकारी|के\\s*बारे\\s*में|बारे\\s*में|शब्द(?:ों)?\\s*में|शब्द(?:ों)?|बताओ|बताइए|बताइये|बताएं|बताएँ|बताओं|समझाओ|समझाइए|समझाइये|लिखो|लिखिए|लिखिये|दीजिए|दीजिये|दो|मुझे|कृपया|एक|पर|का|की|के|को|है|हैं|क्या|कौन|कैसे|में|से)(?=\\s|$|[?।,.])')
_NOTE_HI = 'ध्यान दें: विस्तृत उत्तर नहीं बन सका; नीचे छोटा उत्तर है।\n\n'
_NF_HI = 'हिन्दी विकिपीडिया में इस विषय का लेख नहीं मिला, इसलिए मैं अनुमान से उत्तर नहीं दे रहा।'

def _on_hi(fn, *a, **k):
    global API
    old = API
    API = HI_API
    try:
        return fn(*a, **k)
    finally:
        API = old

def _hi_is_prose(p):
    q = re.sub(r"\([^()]*\)", "", p)
    return len(p) >= 120 and "[...]" not in p and q.count(":") <= 4 and not (q.count(",") >= 6 and len(q) / (q.count(",") + 1) < 20)

def _hi_back(s):
    h = " ".join(s.split()[:4]) + " "
    return h.startswith(_HI_BACK) or any(w in h for w in ("उसे ", "उसके ", "उन्हें ", "उनके "))

def _hi_sents_raw(p):
    out = []
    for s in re.split(r"(?<=[।?!])\s+", p.strip()):
        s = s.strip()
        if 40 <= len(s) <= 600 and re.search(r"[।?!]$", s) and s.count("(") == s.count(")") and s.count("[") == s.count("]") and "[...]" not in s:
            out.append(s)
    return out

def _hi_sents(p):
    return [x for x in _hi_sents_raw(p) if not _hi_back(x)]

def _hi_tidy(s):
    for _ in range(3):
        s = re.sub(r"\([^()]*\)", lambda m: "" if any(c.isalpha() and not ("\u0900" <= c <= "\u097F") for c in m.group(0)) else m.group(0), s)
    s = re.sub(r"\(\s*[^()\s]*\s*[:：]?\s*\)", "", s)
    s = re.sub(r"\s+([,।;:?!)])", r"\1", s)
    s = re.sub(r"\(\s+", "(", s)
    return re.sub(r"\s{2,}", " ", s).strip()

def _hi_text(sel, pool, order):
    parts = []
    for n in sorted(sel, key=lambda k: order[k]):
        if not sel[n]:
            continue
        body = " ".join(_hi_tidy(pool[n][i]) for i in sorted(sel[n]))
        parts.append(body if n == "Lead" else n + "\n" + body)
    return "\n\n".join(parts)

def _hi_fit(sel, pool, order, words, allowed=None, lo=0.90, hi=1.10, protect=None):
    def size(s):
        return len(_hi_text(s, pool, order).split())
    cur = size(sel)
    for _ in range(30):
        if cur <= hi * words:
            break
        best = None
        for n in list(sel):
            for i in list(sel[n]):
                if n == "Lead" and len(sel[n]) <= 1:
                    continue
                if protect is not None and n == protect[0] and i == protect[1]:
                    continue
                s2 = {k: list(v) for k, v in sel.items()}
                s2[n].remove(i)
                if not s2[n]:
                    del s2[n]
                z = size(s2)
                if abs(z - words) < abs(cur - words) and (best is None or abs(z - words) < best[0]):
                    best = (abs(z - words), s2, z)
        if not best:
            break
        sel, cur = best[1], best[2]
    for _ in range(14):
        if cur >= lo * words:
            break
        best = None
        for n, sents in pool.items():
            if allowed is not None and n not in allowed:
                continue
            for i in range(len(sents)):
                if n in sel and i in sel[n]:
                    continue
                s2 = {k: list(v) for k, v in sel.items()}
                s2.setdefault(n, []).append(i)
                z = size(s2)
                if z > hi * words:
                    continue
                score = abs(z - words) + (0 if n in sel else 8)
                if abs(z - words) < abs(cur - words) and (best is None or score < best[0]):
                    best = (score, s2, z)
        if not best:
            break
        sel, cur = best[1], best[2]
    return sel, cur

def hi_long_answer_t(topic, title, words=250, seed=0):
    t0 = time.time()
    secs = _on_hi(_la_sections, title)
    fam, order, total = {}, [], {}
    for s in secs[1:]:
        if s[0] in HI_SKIP or s[3] in HI_SKIP:
            continue
        prose = [p for p in s[1] if _hi_is_prose(p)]
        if not prose:
            continue
        total[s[3]] = total.get(s[3], 0) + sum(len(p) for p in prose)
        if s[3] not in fam:
            fam[s[3]] = [s[0], [], len(order)]
            order.append(s[3])
        if len(fam[s[3]][1]) < 2:
            fam[s[3]][1].extend(prose[:2 - len(fam[s[3]][1])])
    names = order[:30]
    lead_pars = [p for p in secs[0][1] if _hi_is_prose(p)][:(2 if names else 8)]
    lead_sents = [x for p in lead_pars for x in _hi_sents(p)]
    if not names and len(lead_sents) < 4:
        return {"ok": False, "text": "इस लेख में पर्याप्त पाठ नहीं मिला।", "why": "no sections"}
    chosen, chosen_idx = [], []
    if names:
        k_sec = max(2, min(7, round(words / 45) - 2))
        numbered = "\n".join(f"{i + 1}. {n[:50]}" for i, n in enumerate(names))
        r = ask_long(f"Topic: {topic}. Here are the section headings (in Hindi) of its encyclopedia article.\n\n{numbered}\n\n"
                     f"Reply ONLY with the numbers of the {k_sec} headings most important for a general overview of the topic "
                     "(prefer the core subject matter; avoid minor, trivia, research-history or list sections). Example reply: 2, 4, 5, 7",
                     seed, n=24).strip()
        for x in re.findall(r"\d+", r):
            v = int(x)
            if 1 <= v <= len(names) and v not in chosen_idx:
                chosen_idx.append(v)
        chosen = [names[v - 1] for v in chosen_idx[:k_sec]]
        for n in sorted(names, key=lambda p: -total[p]):
            if len(chosen) >= k_sec:
                break
            if n not in chosen:
                chosen.append(n)
        chosen.sort(key=lambda n: fam[n][2])
    pool = {"Lead": lead_sents}
    for n in names:
        pool[n] = [x for par in fam[n][1] for x in _hi_sents(par)]
    ord_ = {"Lead": -1}
    ord_.update({n: fam[n][2] for n in names})
    sel = {}
    if lead_sents:
        kk = 3 if names else max(3, round(words / 35))
        lead_keep = _la_pick(topic, "Lead", lead_sents, kk, seed)
        sel["Lead"] = sorted(set([lead_sents.index(x) for x in lead_keep if x in lead_sents] + [0]))
    per = max(1, round(words / max(1, len(chosen)) / 40))
    for n in chosen:
        if pool[n]:
            keep = _la_pick(topic, fam[n][0], pool[n], min(2, per), seed)
            sel[n] = [pool[n].index(x) for x in keep if x in pool[n]]
    sel = {k: v for k, v in sel.items() if v}
    if not sel:
        return {"ok": False, "text": "इस लेख में पर्याप्त पाठ नहीं मिला।", "why": "nothing selected"}
    prot = ("Lead", 0) if "Lead" in sel else None
    sel, cur = _hi_fit(sel, pool, ord_, words, allowed=set(["Lead"] + chosen), protect=prot)
    if cur < 0.90 * words:
        sel, cur = _hi_fit(sel, pool, ord_, words, allowed=None, protect=prot)
    page_text = " ".join(" ".join(s[1]) for s in secs)
    exact = all(pool[n][i] in page_text for n in sel for i in sel[n])
    body = _hi_text(sel, pool, ord_)
    try:
        ed = _on_hi(page_edited, title)[:10]
    except Exception:
        ed = "unknown"
    text = body + f"\n\n(स्रोत: हिन्दी विकिपीडिया, {title}, अंतिम संपादन {ed}। पाठ CC BY-SA के अंतर्गत विकिपीडिया से हू-ब-हू कॉपी किया गया; अतिरिक्त स्पेस और दूसरी लिपियों में लिखे नाम वाले कोष्ठक हटाए गए हैं।)"
    return {"ok": True, "text": text, "title": title, "exact": exact, "seconds": round(time.time() - t0, 1),
            "body_words": len(body.split()), "target": words, "lead_only": not names,
            "used_headings": [n for n in sorted(sel, key=lambda k: ord_[k]) if n != "Lead"]}

def hi_request(q):
    q = q.strip().translate(_DEV)
    if not re.search(r"[\u0900-\u097F]", q) or not _HI_LONG.search(q):
        return None
    m = _HI_WORDS.search(q)
    words = max(60, min(int(m.group(1)) if m else 250, 600))
    t = _HI_WORDS.sub(" ", q)
    t = re.sub(r"\d+", " ", t)
    for _ in range(3):
        t = _HI_NOISE.sub(" ", t)
    t = re.sub(r"[?।,.!\"']", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return (words, t) if len(t) >= 2 else None

def hi_title(topic):
    j = _on_hi(wiki_get, {"action": "query", "format": "json", "titles": topic, "redirects": 1})
    pg = list(j["query"]["pages"].values())[0]
    if "missing" not in pg and "invalid" not in pg and pg.get("title"):
        return pg["title"]
    j = _on_hi(wiki_get, {"action": "query", "format": "json", "list": "search", "srsearch": topic, "srlimit": 8})
    nz = lambda s: re.sub(r"[\s\-\.॰·,()]", "", s)
    tn = nz(topic)
    res = [r["title"] for r in j.get("query", {}).get("search", [])]
    for t in res:
        if tn and nz(t) == tn:
            return t
    for t in res:
        if tn and (tn in nz(t) or nz(t) in tn):
            try:
                j2 = _on_hi(wiki_get, {"action": "query", "format": "json", "titles": t, "redirects": 1})
                return list(j2["query"]["pages"].values())[0].get("title", t)
            except Exception:
                return t
    return None

def agent_long2(q, *a, **k):
    try:
        r = hi_request(q)
    except Exception:
        r = None
    if r:
        words, topic = r
        try:
            title = hi_title(topic)
            if not title:
                print(f"   [no Hindi page found for: {topic}]", flush=True)
                return _NF_HI
            res = hi_long_answer_t(topic, title, words)
            if res.get("ok") and res.get("exact") and res.get("body_words", 0) >= 60:
                return res["text"]
            print(f"   [Hindi long answer not usable: {res.get('why')} exact={res.get('exact')} words={res.get('body_words')}]", flush=True)
        except Exception as e:
            print("   [Hindi long answer failed:", type(e).__name__, str(e)[:80] + "]", flush=True)
        return _NOTE_HI + str(agent_base(q, *a, **k))
    return agent_long(q, *a, **k)



# ======== Telugu long answers ========
# Telugu long answers, experimental (16 Oct 2026). Needs the Hindi file (hi_all_experimental.py) loaded first, because it reuses its helpers.
import re, time
TE_API = 'https://te.wikipedia.org/w/api.php'
TE_SKIP = {'మూలాలు, వనరులు', 'బయటి లంకెలు', 'ఉపయోగించిన మూలాలు', 'ఇతరలింకులు', 'ఇవి కూడా చూడండి', 'బయటి లింకులు', 'ఇతర పఠనాలు', 'వనరులు', 'మూలాలు మరియు వనరులు', 'మూలాలు'}
_TE_BACK1 = {'ఇలా', 'అతను', 'వారు', 'అతని', 'అలా', 'అప్పుడు', 'వీరు', 'దీని', 'అది', 'ఇది', 'దీనికి', 'వాటి', 'ఆ', 'అతడు', 'ఈ', 'అందువల్ల', 'అతన్ని', 'ఆమె', 'ఆయన', 'వీటి', 'ఆ తరువాత', 'అందుకే', 'కాబట్టి', 'దీనిని', 'అయితే', 'ఆమెను', 'తరువాత'}
_TE_BACK4 = ('దీని', 'దీనిని', 'అతని', 'ఆయన', 'వారి', 'వీరి', 'ఆమె ', 'అతడు')

def _on_te(fn, *a, **k):
    global API
    old = API
    API = TE_API
    try:
        return fn(*a, **k)
    finally:
        API = old

def _te_is_prose(p):
    q = re.sub(r"\([^()]*\)", "", p)
    return len(p) >= 120 and "[...]" not in p and q.count(":") <= 4 and not (q.count(",") >= 6 and len(q) / (q.count(",") + 1) < 20)

def _te_sents_raw(p):
    out = []
    for s in re.split(r"(?<=[\.?!।])\s+", p.strip()):
        s = s.strip()
        if 40 <= len(s) <= 600 and re.search(r"[\.?!।]$", s) and s.count("(") == s.count(")") and s.count("[") == s.count("]") and "[...]" not in s and re.search(r"[\u0C00-\u0C7F]", s):
            out.append(s)
    return out

def _te_back(s):
    w = s.split()
    if not w:
        return True
    return w[0] in _TE_BACK1 or any(x in _TE_BACK4 for x in w[:4])

def _te_sents(p):
    return [x for x in _te_sents_raw(p) if not _te_back(x)]

def _te_tidy(s):
    s = re.sub(r"\(\s*వినండి\s*\)", "", s)
    for _ in range(3):
        s = re.sub(r"\([^()]*\)", lambda m: "" if any(c.isalpha() and not ("\u0C00" <= c <= "\u0C7F") for c in m.group(0)) else m.group(0), s)
    s = re.sub(r"\(\s*[^()\s]*\s*[:：]?\s*\)", "", s)
    s = re.sub(r"\s+([,।;:?!).])", r"\1", s)
    s = re.sub(r"\(\s+", "(", s)
    return re.sub(r"\s{2,}", " ", s).strip()

def te_title(en_topic):
    j = wiki_get({"action": "query", "format": "json", "titles": en_topic, "prop": "langlinks", "lllang": "te", "redirects": 1})
    pg = list(j["query"]["pages"].values())[0]
    tl = pg.get("langlinks", [])
    return tl[0]["*"] if tl else None

def _te_text(sel, pool, order):
    parts = []
    for n in sorted(sel, key=lambda k: order[k]):
        if not sel[n]:
            continue
        body = " ".join(_te_tidy(pool[n][i]) for i in sorted(sel[n]))
        parts.append(body if n == "Lead" else n + "\n" + body)
    return "\n\n".join(parts)

def _te_fit(sel, pool, order, words, allowed=None, lo=0.90, hi=1.10, protect=None):
    def size(s):
        return len(_te_text(s, pool, order).split())
    cur = size(sel)
    for _ in range(30):
        if cur <= hi * words:
            break
        best = None
        for n in list(sel):
            for i in list(sel[n]):
                if n == "Lead" and len(sel[n]) <= 1:
                    continue
                if protect is not None and n == protect[0] and i == protect[1]:
                    continue
                s2 = {k: list(v) for k, v in sel.items()}
                s2[n].remove(i)
                if not s2[n]:
                    del s2[n]
                z = size(s2)
                if abs(z - words) < abs(cur - words) and (best is None or abs(z - words) < best[0]):
                    best = (abs(z - words), s2, z)
        if not best:
            break
        sel, cur = best[1], best[2]
    for _ in range(14):
        if cur >= lo * words:
            break
        best = None
        for n, sents in pool.items():
            if allowed is not None and n not in allowed:
                continue
            for i in range(len(sents)):
                if n in sel and i in sel[n]:
                    continue
                s2 = {k: list(v) for k, v in sel.items()}
                s2.setdefault(n, []).append(i)
                z = size(s2)
                if z > hi * words:
                    continue
                score = abs(z - words) + (0 if n in sel else 8)
                if abs(z - words) < abs(cur - words) and (best is None or score < best[0]):
                    best = (score, s2, z)
        if not best:
            break
        sel, cur = best[1], best[2]
    return sel, cur

def te_long_answer_t(topic, title, words=250, seed=0):
    t0 = time.time()
    secs = _on_te(_la_sections, title)
    fam, order, total = {}, [], {}
    for s in secs[1:]:
        if s[0] in TE_SKIP or s[3] in TE_SKIP:
            continue
        prose = [p for p in s[1] if _te_is_prose(p)]
        if not prose:
            continue
        total[s[3]] = total.get(s[3], 0) + sum(len(p) for p in prose)
        if s[3] not in fam:
            fam[s[3]] = [s[0], [], len(order)]
            order.append(s[3])
        if len(fam[s[3]][1]) < 2:
            fam[s[3]][1].extend(prose[:2 - len(fam[s[3]][1])])
    names = order[:30]
    lead_pars = [p for p in secs[0][1] if _te_is_prose(p)][:(2 if names else 8)]
    lead_sents = [x for p in lead_pars for x in _te_sents(p)]
    if not names and len(lead_sents) < 4:
        return {"ok": False, "text": "ఈ వ్యాసంలో తగినంత వచనం దొరకలేదు.", "why": "no sections"}
    chosen, chosen_idx = [], []
    if names:
        k_sec = max(2, min(7, round(words / 45) - 2))
        numbered = "\n".join(f"{i + 1}. {n[:50]}" for i, n in enumerate(names))
        r = ask_long(f"Topic: {topic}. Here are the section headings (in Telugu) of its encyclopedia article.\n\n{numbered}\n\n"
                     f"Reply ONLY with the numbers of the {k_sec} headings most important for a general overview of the topic "
                     "(prefer the core subject matter; avoid minor, trivia, research-history or list sections). Example reply: 2, 4, 5, 7",
                     seed, n=24).strip()
        for x in re.findall(r"\d+", r):
            v = int(x)
            if 1 <= v <= len(names) and v not in chosen_idx:
                chosen_idx.append(v)
        chosen = [names[v - 1] for v in chosen_idx[:k_sec]]
        for n in sorted(names, key=lambda p: -total[p]):
            if len(chosen) >= k_sec:
                break
            if n not in chosen:
                chosen.append(n)
        chosen.sort(key=lambda n: fam[n][2])
    pool = {"Lead": lead_sents}
    for n in names:
        pool[n] = [x for par in fam[n][1] for x in _te_sents(par)]
    ord_ = {"Lead": -1}
    ord_.update({n: fam[n][2] for n in names})
    sel = {}
    if lead_sents:
        kk = 3 if names else max(3, round(words / 35))
        lead_keep = _la_pick(topic, "Lead", lead_sents, kk, seed)
        sel["Lead"] = sorted(set([lead_sents.index(x) for x in lead_keep if x in lead_sents] + [0]))
    per = max(1, round(words / max(1, len(chosen)) / 40))
    for n in chosen:
        if pool[n]:
            keep = _la_pick(topic, fam[n][0], pool[n], min(2, per), seed)
            sel[n] = [pool[n].index(x) for x in keep if x in pool[n]]
    sel = {k: v for k, v in sel.items() if v}
    if not sel:
        return {"ok": False, "text": "ఈ వ్యాసంలో తగినంత వచనం దొరకలేదు.", "why": "nothing selected"}
    prot = ("Lead", 0) if "Lead" in sel else None
    sel, cur = _te_fit(sel, pool, ord_, words, allowed=set(["Lead"] + chosen), protect=prot)
    if cur < 0.90 * words:
        sel, cur = _te_fit(sel, pool, ord_, words, allowed=None, protect=prot)
    page_text = " ".join(" ".join(s[1]) for s in secs)
    exact = all(pool[n][i] in page_text for n in sel for i in sel[n])
    body = _te_text(sel, pool, ord_)
    try:
        ed = _on_te(page_edited, title)[:10]
    except Exception:
        ed = "unknown"
    text = body + f"\n\n(మూలం: తెలుగు వికీపీడియా, {title}, చివరి మార్పు {ed}. వచనం CC BY-SA కింద వికీపీడియా నుండి యథాతథంగా కాపీ చేయబడింది; అదనపు ఖాళీలు, ఇతర లిపులలో పేర్లున్న కుండలీకరణాలు తొలగించబడ్డాయి.)"
    return {"ok": True, "text": text, "title": title, "exact": exact, "seconds": round(time.time() - t0, 1),
            "body_words": len(body.split()), "target": words, "lead_only": not names,
            "used_headings": [n for n in sorted(sel, key=lambda k: ord_[k]) if n != "Lead"]}



# ======== Telugu wiring ========
# Telugu wiring + one entry point agent_all (experimental, 16 Oct 2026). Needs hi_all_experimental.py and te_all_experimental.py loaded first.
import re, time
_TE_DIG = {3174: 48, 3175: 49, 3176: 50, 3177: 51, 3178: 52, 3179: 53, 3180: 54, 3181: 55, 3182: 56, 3183: 57}
_TE_TRIG = ('వివరంగా', 'విపులంగా', 'వివరణాత్మకంగా', 'వ్యాసం', 'వ్యాసము', 'పూర్తి సమాచారం', 'సమగ్ర సమాచారం', 'పదాలలో', 'పదాల్లో', 'పదాల')
_TE_NOISE = {'సమాచారం', 'నాకు', 'పదాల్లో', 'రాయండి', 'వ్యాసము', 'వ్యాసం', 'వివరించండి', 'పదాల', 'ఇవ్వండి', 'చెప్పు', 'రాయి', 'వివరించు', 'తెలియజేయండి', 'పదాలలో', 'దయచేసి', 'సమగ్ర', 'పై', 'రాయండీ', 'పూర్తి', 'మీద', 'చెప్పండి', 'తెలుపండి', 'గురించి', 'వివరంగా', 'వివరణాత్మకంగా', 'ఇవ్వు', 'గూర్చి', 'విపులంగా', 'ఒక'}
_NF_TE = 'తెలుగు వికీపీడియాలో ఈ అంశం దొరకలేదు, కాబట్టి ఊహించి సమాధానం ఇవ్వను.'
_NOTE_TE = 'గమనిక: వివరమైన సమాధానం తయారు కాలేదు; క్రింద చిన్న సమాధానం ఉంది.\n\n'

def _te_tidy(s):
    s = s.replace("|", " ")
    s = re.sub(r"\(\s*వినండి\s*\)", "", s)
    for _ in range(3):
        s = re.sub(r"\([^()]*\)", lambda m: "" if any(c.isalpha() and not ("\u0C00" <= c <= "\u0C7F") for c in m.group(0)) else m.group(0), s)
    s = re.sub(r"\(\s*[^()\s]*\s*[:：]?\s*\)", "", s)
    s = re.sub(r"\s+([,।;:?!).])", r"\1", s)
    s = re.sub(r"\(\s+", "(", s)
    return re.sub(r"\s{2,}", " ", s).strip()

def te_request(q):
    q = q.strip().translate(_TE_DIG)
    if not re.search(r"[\u0C00-\u0C7F]", q) or not any(t in q for t in _TE_TRIG):
        return None
    m = re.search(r"(\d+)\s*(?:పదాల|పదముల)", q)
    words = max(60, min(int(m.group(1)) if m else 250, 600))
    toks = [t for t in re.sub(r"[?.,!\"']", " ", q).split() if t not in _TE_NOISE and not re.fullmatch(r"\d+", t)]
    topic = " ".join(toks).strip()
    return (words, topic) if len(topic) >= 2 else None

def te_find_title(topic):
    j = _on_te(wiki_get, {"action": "query", "format": "json", "titles": topic, "redirects": 1})
    pg = list(j["query"]["pages"].values())[0]
    if "missing" not in pg and "invalid" not in pg and pg.get("title"):
        return pg["title"]
    j = _on_te(wiki_get, {"action": "query", "format": "json", "list": "search", "srsearch": topic, "srlimit": 8})
    nz = lambda s: re.sub(r"[\s\-\.,()]", "", s)
    tn = nz(topic)
    stem = re.sub(r"[\u0C3E-\u0C4D]+$", "", tn)
    res = [r["title"] for r in j.get("query", {}).get("search", [])]
    for t in res:
        if tn and nz(t) == tn:
            return t
    for t in res:
        if tn and (tn in nz(t) or nz(t) in tn or (len(stem) >= 3 and stem in nz(t))):
            try:
                j2 = _on_te(wiki_get, {"action": "query", "format": "json", "titles": t, "redirects": 1})
                return list(j2["query"]["pages"].values())[0].get("title", t)
            except Exception:
                return t
    return None

def agent_all(q, *a, **k):
    if re.search(r"[\u0C00-\u0C7F]", q):
        try:
            r = te_request(q)
        except Exception:
            r = None
        if r:
            words, topic = r
            try:
                title = te_find_title(topic)
                if not title:
                    print(f"   [no Telugu page found for: {topic}]", flush=True)
                    return _NF_TE
                res = te_long_answer_t(topic, title, words)
                if res.get("ok") and res.get("exact") and res.get("body_words", 0) >= 60:
                    return res["text"]
                print(f"   [Telugu long answer not usable: {res.get('why')} exact={res.get('exact')} words={res.get('body_words')}]", flush=True)
            except Exception as e:
                print("   [Telugu long answer failed:", type(e).__name__, str(e)[:80] + "]", flush=True)
            return _NOTE_TE + str(agent_base(q, *a, **k))
        return agent_base(q, *a, **k)
    return agent_long2(q, *a, **k)



# ======== save guard ========
# Save-guard patch (experimental, 16 Oct 2026). Load after agent_v27d_kaggle.py. Stops answers that say the sources disagree from being saved to kb.json.
_kb_good_v0 = _kb_good

def _kb_good(text):
    if any(v in text for v in DISAGREE.values()):
        return False
    return _kb_good_v0(text)



# ======== fixes: emoji filter ========
# Fixes (16 Oct 2026): emoji/symbol sentence filter and 'use more paragraphs when short' retry. Load after hi_all, te_all and te_wiring files. Needs TE_SRC (built by te_all) - see cell text below.
_SYM = re.compile(r"[\U0001F300-\U0001FAFF\u2600-\u27BF\u25A0-\u25FF\u2022\u25BA\u2B50\uFE0F]")
_hi_sents_raw_v0 = _hi_sents_raw
def _hi_sents_raw(p):
    return [x for x in _hi_sents_raw_v0(p) if not _SYM.search(x)]
_te_sents_raw_v0 = _te_sents_raw
def _te_sents_raw(p):
    return [x for x in _te_sents_raw_v0(p) if not _SYM.search(x)]


# ======== fixes: more paragraphs when short (applied at build time) ========
def hi_long_answer_t1(topic, title, words=250, seed=0, maxp=2):
    t0 = time.time()
    secs = _on_hi(_la_sections, title)
    fam, order, total = {}, [], {}
    for s in secs[1:]:
        if s[0] in HI_SKIP or s[3] in HI_SKIP:
            continue
        prose = [p for p in s[1] if _hi_is_prose(p)]
        if not prose:
            continue
        total[s[3]] = total.get(s[3], 0) + sum(len(p) for p in prose)
        if s[3] not in fam:
            fam[s[3]] = [s[0], [], len(order)]
            order.append(s[3])
        if len(fam[s[3]][1]) < maxp:
            fam[s[3]][1].extend(prose[:maxp - len(fam[s[3]][1])])
    names = order[:30]
    lead_pars = [p for p in secs[0][1] if _hi_is_prose(p)][:(2 if names else 8)]
    lead_sents = [x for p in lead_pars for x in _hi_sents(p)]
    if not names and len(lead_sents) < 4:
        return {"ok": False, "text": "इस लेख में पर्याप्त पाठ नहीं मिला।", "why": "no sections"}
    chosen, chosen_idx = [], []
    if names:
        k_sec = max(2, min(7, round(words / 45) - 2))
        numbered = "\n".join(f"{i + 1}. {n[:50]}" for i, n in enumerate(names))
        r = ask_long(f"Topic: {topic}. Here are the section headings (in Hindi) of its encyclopedia article.\n\n{numbered}\n\n"
                     f"Reply ONLY with the numbers of the {k_sec} headings most important for a general overview of the topic "
                     "(prefer the core subject matter; avoid minor, trivia, research-history or list sections). Example reply: 2, 4, 5, 7",
                     seed, n=24).strip()
        for x in re.findall(r"\d+", r):
            v = int(x)
            if 1 <= v <= len(names) and v not in chosen_idx:
                chosen_idx.append(v)
        chosen = [names[v - 1] for v in chosen_idx[:k_sec]]
        for n in sorted(names, key=lambda p: -total[p]):
            if len(chosen) >= k_sec:
                break
            if n not in chosen:
                chosen.append(n)
        chosen.sort(key=lambda n: fam[n][2])
    pool = {"Lead": lead_sents}
    for n in names:
        pool[n] = [x for par in fam[n][1] for x in _hi_sents(par)]
    ord_ = {"Lead": -1}
    ord_.update({n: fam[n][2] for n in names})
    sel = {}
    if lead_sents:
        kk = 3 if names else max(3, round(words / 35))
        lead_keep = _la_pick(topic, "Lead", lead_sents, kk, seed)
        sel["Lead"] = sorted(set([lead_sents.index(x) for x in lead_keep if x in lead_sents] + [0]))
    per = max(1, round(words / max(1, len(chosen)) / 40))
    for n in chosen:
        if pool[n]:
            keep = _la_pick(topic, fam[n][0], pool[n], min(2, per), seed)
            sel[n] = [pool[n].index(x) for x in keep if x in pool[n]]
    sel = {k: v for k, v in sel.items() if v}
    if not sel:
        return {"ok": False, "text": "इस लेख में पर्याप्त पाठ नहीं मिला।", "why": "nothing selected"}
    prot = ("Lead", 0) if "Lead" in sel else None
    sel, cur = _hi_fit(sel, pool, ord_, words, allowed=set(["Lead"] + chosen), protect=prot)
    if cur < 0.90 * words:
        sel, cur = _hi_fit(sel, pool, ord_, words, allowed=None, protect=prot)
    page_text = " ".join(" ".join(s[1]) for s in secs)
    exact = all(pool[n][i] in page_text for n in sel for i in sel[n])
    body = _hi_text(sel, pool, ord_)
    try:
        ed = _on_hi(page_edited, title)[:10]
    except Exception:
        ed = "unknown"
    text = body + f"\n\n(स्रोत: हिन्दी विकिपीडिया, {title}, अंतिम संपादन {ed}। पाठ CC BY-SA के अंतर्गत विकिपीडिया से हू-ब-हू कॉपी किया गया; अतिरिक्त स्पेस और दूसरी लिपियों में लिखे नाम वाले कोष्ठक हटाए गए हैं।)"
    return {"ok": True, "text": text, "title": title, "exact": exact, "seconds": round(time.time() - t0, 1),
            "body_words": len(body.split()), "target": words, "lead_only": not names,
            "used_headings": [n for n in sorted(sel, key=lambda k: ord_[k]) if n != "Lead"]}


def te_long_answer_t1(topic, title, words=250, seed=0, maxp=2):
    t0 = time.time()
    secs = _on_te(_la_sections, title)
    fam, order, total = {}, [], {}
    for s in secs[1:]:
        if s[0] in TE_SKIP or s[3] in TE_SKIP:
            continue
        prose = [p for p in s[1] if _te_is_prose(p)]
        if not prose:
            continue
        total[s[3]] = total.get(s[3], 0) + sum(len(p) for p in prose)
        if s[3] not in fam:
            fam[s[3]] = [s[0], [], len(order)]
            order.append(s[3])
        if len(fam[s[3]][1]) < maxp:
            fam[s[3]][1].extend(prose[:maxp - len(fam[s[3]][1])])
    names = order[:30]
    lead_pars = [p for p in secs[0][1] if _te_is_prose(p)][:(2 if names else 8)]
    lead_sents = [x for p in lead_pars for x in _te_sents(p)]
    if not names and len(lead_sents) < 4:
        return {"ok": False, "text": "ఈ వ్యాసంలో తగినంత వచనం దొరకలేదు.", "why": "no sections"}
    chosen, chosen_idx = [], []
    if names:
        k_sec = max(2, min(7, round(words / 45) - 2))
        numbered = "\n".join(f"{i + 1}. {n[:50]}" for i, n in enumerate(names))
        r = ask_long(f"Topic: {topic}. Here are the section headings (in Telugu) of its encyclopedia article.\n\n{numbered}\n\n"
                     f"Reply ONLY with the numbers of the {k_sec} headings most important for a general overview of the topic "
                     "(prefer the core subject matter; avoid minor, trivia, research-history or list sections). Example reply: 2, 4, 5, 7",
                     seed, n=24).strip()
        for x in re.findall(r"\d+", r):
            v = int(x)
            if 1 <= v <= len(names) and v not in chosen_idx:
                chosen_idx.append(v)
        chosen = [names[v - 1] for v in chosen_idx[:k_sec]]
        for n in sorted(names, key=lambda p: -total[p]):
            if len(chosen) >= k_sec:
                break
            if n not in chosen:
                chosen.append(n)
        chosen.sort(key=lambda n: fam[n][2])
    pool = {"Lead": lead_sents}
    for n in names:
        pool[n] = [x for par in fam[n][1] for x in _te_sents(par)]
    ord_ = {"Lead": -1}
    ord_.update({n: fam[n][2] for n in names})
    sel = {}
    if lead_sents:
        kk = 3 if names else max(3, round(words / 35))
        lead_keep = _la_pick(topic, "Lead", lead_sents, kk, seed)
        sel["Lead"] = sorted(set([lead_sents.index(x) for x in lead_keep if x in lead_sents] + [0]))
    per = max(1, round(words / max(1, len(chosen)) / 40))
    for n in chosen:
        if pool[n]:
            keep = _la_pick(topic, fam[n][0], pool[n], min(2, per), seed)
            sel[n] = [pool[n].index(x) for x in keep if x in pool[n]]
    sel = {k: v for k, v in sel.items() if v}
    if not sel:
        return {"ok": False, "text": "ఈ వ్యాసంలో తగినంత వచనం దొరకలేదు.", "why": "nothing selected"}
    prot = ("Lead", 0) if "Lead" in sel else None
    sel, cur = _te_fit(sel, pool, ord_, words, allowed=set(["Lead"] + chosen), protect=prot)
    if cur < 0.90 * words:
        sel, cur = _te_fit(sel, pool, ord_, words, allowed=None, protect=prot)
    page_text = " ".join(" ".join(s[1]) for s in secs)
    exact = all(pool[n][i] in page_text for n in sel for i in sel[n])
    body = _te_text(sel, pool, ord_)
    try:
        ed = _on_te(page_edited, title)[:10]
    except Exception:
        ed = "unknown"
    text = body + f"\n\n(మూలం: తెలుగు వికీపీడియా, {title}, చివరి మార్పు {ed}. వచనం CC BY-SA కింద వికీపీడియా నుండి యథాతథంగా కాపీ చేయబడింది; అదనపు ఖాళీలు, ఇతర లిపులలో పేర్లున్న కుండలీకరణాలు తొలగించబడ్డాయి.)"
    return {"ok": True, "text": text, "title": title, "exact": exact, "seconds": round(time.time() - t0, 1),
            "body_words": len(body.split()), "target": words, "lead_only": not names,
            "used_headings": [n for n in sorted(sel, key=lambda k: ord_[k]) if n != "Lead"]}


def _wrap_retry(f1):
    def g(topic, title, words=250, seed=0):
        r = f1(topic, title, words, seed, 2)
        if (not r.get("ok")) or r.get("body_words", 0) < 0.9 * words:
            r2 = f1(topic, title, words, seed, 6)
            if r2.get("ok") and ((not r.get("ok")) or abs(r2["body_words"] - words) < abs(r["body_words"] - words)):
                r2["more_text_used"] = True
                return r2
        return r
    return g
hi_long_answer_t = _wrap_retry(hi_long_answer_t1)
te_long_answer_t = _wrap_retry(te_long_answer_t1)


# ======== thin-page notes ========
# Thin-page notes (16 Oct 2026): agent_long2 and agent_all return the exact short text with a note when the article is short. Load after the other Hindi and Telugu files.
_THIN_HI = "\n(ध्यान दें: हिन्दी विकिपीडिया पर इस विषय का लेख छोटा है, इसलिए माँगे गए शब्दों से कम पाठ ही उपलब्ध है।)"
_THIN_TE = "\n(గమనిక: తెలుగు వికీపీడియాలో ఈ వ్యాసం చిన్నది, అందుకే కోరినన్ని పదాలు అందుబాటులో లేవు.)"

def _usable(res, words):
    return res.get("ok") and res.get("exact") and res.get("body_words", 0) >= 25

def _thin(res, words, note):
    return note if res.get("body_words", 0) < 0.85 * words else ""

def agent_long2(q, *a, **k):
    try:
        r = hi_request(q)
    except Exception:
        r = None
    if r:
        words, topic = r
        try:
            title = hi_title(topic)
            if not title:
                print(f"   [no Hindi page found for: {topic}]", flush=True)
                return _NF_HI
            res = hi_long_answer_t(topic, title, words)
            if _usable(res, words):
                return res["text"] + _thin(res, words, _THIN_HI)
            print(f"   [Hindi long answer not usable: {res.get('why')} exact={res.get('exact')} words={res.get('body_words')}]", flush=True)
        except Exception as e:
            print("   [Hindi long answer failed:", type(e).__name__, str(e)[:80] + "]", flush=True)
        return _NOTE_HI + str(agent_base(q, *a, **k))
    return agent_long(q, *a, **k)

def agent_all(q, *a, **k):
    if re.search(r"[\u0C00-\u0C7F]", q):
        try:
            r = te_request(q)
        except Exception:
            r = None
        if r:
            words, topic = r
            try:
                title = te_find_title(topic)
                if not title:
                    print(f"   [no Telugu page found for: {topic}]", flush=True)
                    return _NF_TE
                res = te_long_answer_t(topic, title, words)
                if _usable(res, words):
                    return res["text"] + _thin(res, words, _THIN_TE)
                print(f"   [Telugu long answer not usable: {res.get('why')} exact={res.get('exact')} words={res.get('body_words')}]", flush=True)
            except Exception as e:
                print("   [Telugu long answer failed:", type(e).__name__, str(e)[:80] + "]", flush=True)
            return _NOTE_TE + str(agent_base(q, *a, **k))
        return agent_base(q, *a, **k)
    return agent_long2(q, *a, **k)



# ======== EN request parser v6 ========
# English request parser v6 (16 Oct 2026): strips filler words from topics, routes 'how X works in detail' to the long answer. Load after long_answer_v5_5_patch.
if "_la_request_v5" not in globals():
    _la_request_v5 = la_request

_LA_LEADWORD = re.compile(r"^(?:give me|give|summari[sz]e|describe|explain|explanation of|description of|overview of|account of|a|an|the|about|me)\s+", re.I)
def la_request(q):
    r = _la_request_v5(q)
    if r:
        words, t = r
        t = t.strip()
        for _ in range(4):
            t2 = _LA_LEADWORD.sub("", t)
            if t2 == t:
                break
            t = t2
        return (words, t) if len(t) >= 2 else None
    if re.search(r"\bin[- ]depth\b|\bin detail\b", q, re.I):
        m = re.search(r"\bhow\s+(?:does\s+|do\s+)?(?:a |an |the )?(.+?)\s+works?\b", q, re.I)
        if m and 2 <= len(m.group(1)) <= 40:
            return (250, m.group(1).strip())
    return None



# ======== v28b (17 Oct 2026): pinned-only alias lookup (what/which merged, possessive 's dropped) ========
def new_key(eng):
    t = re.sub(r"['\u2019]s\b", "", eng.lower())
    words = ["what" if w == "which" else w for w in re.findall(r"[a-z0-9]+", t)]
    return "|".join(sorted(set(w for w in words if w in _QW or w not in _STOP)))

if not globals().get("_alias_ready"):
    _kb_get_prev_alias = kb_get
    def kb_get(db, eng, lang, now=None):
        e = _kb_get_prev_alias(db, eng, lang, now)
        if e:
            return e
        nk = new_key(eng)
        if not nk:
            return None
        for v in db.values():
            if v.get("pinned") and v.get("lang") == lang and new_key(v.get("q", "")) == nk:
                return v
        return None
    _alias_ready = True


# ======== v28c (17 Oct 2026): guard patches - do not save code/refusal answers; wider code-task detection ========
import re as _re_g
_CODE_TASK = _re_g.compile(
    r"\b(write|create|make|build|generate|fix|debug|run|code)\b.{0,80}\b(python|code|function|program|script|csv|json|excel|html|javascript|sql)\b"
    r"|\bpython code\b|\b(csv|json|xlsx|docx)\b.{0,30}\bfile\b", _re_g.I)
_BAD_PHRASES = ("As a large language model", "I do not have direct access", "I cannot access", "I'm unable to",
                "I am unable to", "As an AI")
if not globals().get("_guard2_ready"):
    _kb_skip_prev2 = _kb_skip
    def _kb_skip(question):
        if _kb_skip_prev2(question):
            return True
        return bool(_CODE_TASK.search(question))
    _kb_good_prev2 = _kb_good
    def _kb_good(text):
        if "```" in text or any(p in text for p in _BAD_PHRASES):
            return False
        return _kb_good_prev2(text)
    _guard2_ready = True
if not globals().get("_guard3_ready"):
    _kb_skip_prev3 = _kb_skip
    def _kb_skip(question):
        if _re_g.search(r"\bprogramming\b", question, _re_g.I) and not _CODE_TASK.search(question):
            return False
        return _kb_skip_prev3(question)
    _guard3_ready = True


# ======== v28d (18 Oct 2026): code-request router (verb + code word) placed in front of agent and agent_base ========
import re as _re_r
_IMPERATIVE = _re_r.compile(r"(^|[.;:!?]\s+)(please\s+|kindly\s+)?(can you\s+|could you\s+|would you\s+)?"
                            r"(write|create|make|build|generate|fix|debug|run|code)\b", _re_r.I)
_CODEWORD = _re_r.compile(r"\b(python|code|function|program|script|csv|json|excel|html|javascript|sql|xlsx|docx)\b", _re_r.I)
def is_code_request(q):
    return bool(_IMPERATIVE.search(q)) and bool(_CODEWORD.search(q))

def _make_code_router(prev):
    def routed(question, *a, **k):
        if is_code_request(question):
            seed = a[0] if a and isinstance(a[0], int) else k.get("seed", 0)
            r = code_task(question, seed)
            if r["ok"]:
                return ("Result:\n" + r["output"] + "\n\nCode used:\n" + r["code"] +
                        f"\n(Ran the code: no errors, tries: {r['tries']}. Running without errors does not prove "
                        "the result is what you wanted.)")
            return "I could not make a working program in 3 tries. Last output:\n" + r["output"]
        return prev(question, *a, **k)
    return routed

if not globals().get("_code_router_ready"):
    agent_base = _make_code_router(agent_base)
    agent = _make_code_router(agent)
    _code_router_ready = True


# ======== v28e (18 Oct 2026): Hindi and Telugu code-request router (exact command verb + code word, whole tokens) ========
import re as _re_i
_TOK = _re_i.compile(r"[^\s\u0964\u0965,.;:!?\"'()\[\]{}<>=+*/\\|`~@#$%^&\-]+")
def _toks(q):
    return [t.replace("\u200c", "").replace("\u200d", "").lower() for t in _TOK.findall(q)]
_DEV, _TEL = _re_i.compile(r"[\u0900-\u097F]"), _re_i.compile(r"[\u0C00-\u0C7F]")
_HI_VERBS = ["लिखो", "लिखिए", "लिखें", "बनाओ", "बनाइए", "बनाएं", "चलाओ", "चलाइए", "चलाएं",
             "ठीक करो", "ठीक कीजिए", "जनरेट करो"]
_TE_VERBS = ["రాయండి", "రాయి", "తయారు చేయండి", "చేయండి", "నడపండి"]
_IND_CW = ["कोड", "प्रोग्राम", "पायथन", "स्क्रिप्ट", "फंक्शन", "एचटीएमएल",
           "కోడ్", "ప్రోగ్రామ్", "పైథాన్", "స్క్రిప్ట్", "ఫంక్షన్"]
_ENG_CW = {"python", "code", "function", "program", "script", "csv", "json", "excel", "html",
           "javascript", "sql", "xlsx", "docx"}
def _indic_code(q, verbs, script_re):
    if not script_re.search(q):
        return False
    toks = _toks(q)
    s = " " + " ".join(toks) + " "
    has_verb = any((" " + v + " ") in s for v in verbs)
    has_cw = any(t in _ENG_CW or any(t.startswith(w) for w in _IND_CW) for t in toks)
    return has_verb and has_cw
if "_is_code_en" not in globals():
    _is_code_en = is_code_request
def is_code_request(q):
    return _is_code_en(q) or _indic_code(q, _HI_VERBS, _DEV) or _indic_code(q, _TE_VERBS, _TEL)


# ======== v28f (18 Oct 2026): file-compute router (question names an existing file and asks for a count/total/average -> run code on it) ========
import os as _os_f, re as _re_f
_COMPUTE = _re_f.compile(r"\b(how many|total|sum|average|count|number of|mean of)\b", _re_f.I)
def file_compute_path(q):
    m = _FILE_RE.search(q)
    if m and _os_f.path.exists(m.group(0)) and _COMPUTE.search(q):
        return m.group(0)
    return None

def _preview(path, n=6, cap=400):
    try:
        return "\n".join(read_any(path).split("\n")[:n])[:cap]
    except Exception:
        return ""

def _make_file_router3(prev):
    def routed(question, *a, **k):
        p = file_compute_path(question)
        if p:
            ap = _os_f.path.abspath(p)
            seed = a[0] if a and isinstance(a[0], int) else k.get("seed", 0)
            task = ("Write a Python program for this: " + question + f" The file is at {ap}. Open it using exactly that path. "
                    "The first lines of the file are shown below. Use the real column names and positions you see here.\n"
                    "FILE START\n" + _preview(p) + "\nFILE END\n"
                    "Print the final answer on the screen as a number or short text. Do not save or write any files.")
            r = code_task(task, seed)
            if r["ok"]:
                return ("Result:\n" + r["output"] + "\n\nCode used:\n" + r["code"] +
                        f"\n(Ran the code on file {_os_f.path.basename(p)}: no errors, tries: {r['tries']}. "
                        "Running without errors does not prove the result is what you wanted.)")
            return "I could not make a working program in 3 tries. Last output:\n" + r["output"]
        return prev(question, *a, **k)
    return routed



# ======== v28g (18 Oct 2026): file-compute questions run 3 times; answer only if at least 2 runs agree ========
import collections as _col_g, os as _os_g, re as _re_g2
def last_num(s):
    xs = _re_g2.findall(r"\d+(?:\.\d+)?", s.replace(",", ""))
    return float(xs[-1]) if xs else None

def decide_consensus(values):
    got = [v for v in values if v is not None]
    if not got:
        return None, "no run produced a number"
    def near(a, b):
        return abs(a - b) <= 0.0051 + 1e-9 * max(abs(a), abs(b))
    best = []
    for v in got:
        cl = [w for w in got if near(v, w)]
        if len(cl) > len(best):
            best = cl
    if len(best) >= 2:
        ans = _col_g.Counter(best).most_common(1)[0][0]
        return ans, f"{len(best)} of {len(values)} runs agreed"
    return None, "the runs disagreed: " + ", ".join(str(x) for x in got)

def _make_file_router4(prev):
    def routed(question, *a, **k):
        p = file_compute_path(question)
        if not p:
            return prev(question, *a, **k)
        ap = _os_g.path.abspath(p)
        base = a[0] if a and isinstance(a[0], int) else k.get("seed", 0)
        task = ("Write a Python program for this: " + question + f" The file is at {ap}. Open it using exactly that path. "
                "The first lines of the file are shown below. Use the real column names and positions you see here.\n"
                "FILE START\n" + _preview(p) + "\nFILE END\n"
                "Print the final answer on the screen as a number or short text. Do not save or write any files.")
        runs = []
        for s in (base, base + 10, base + 20):
            try:
                r = code_task(task, s)
            except BaseException:
                r = {"ok": False, "output": "", "code": "", "tries": 0}
            v = last_num(r["output"]) if r.get("ok") else None
            runs.append((v, r.get("output", ""), r.get("code", "")))
        ans, status = decide_consensus([x[0] for x in runs])
        if ans is None:
            vals = ", ".join("no number" if x[0] is None else str(x[0]) for x in runs)
            return (f"I ran the program 3 times on {_os_g.path.basename(p)} and got different results ({vals}). "
                    f"I cannot give one answer ({status}). Please check the file or ask in another way.")
        pick = next(x for x in runs if x[0] is not None and abs(x[0] - ans) <= 0.0051 + 1e-9 * abs(ans))
        return ("Result:\n" + pick[1].strip() + "\n\nCode used:\n" + pick[2] +
                f"\n(Ran the code on file {_os_g.path.basename(p)} 3 times: {status}. "
                "Agreement does not prove the result is what you wanted.)")
    return routed

_file_router3_ready = True   # the single-run router of v28f is replaced by the 3-run router below
if not globals().get("_file_router4_ready"):
    agent_base = _make_file_router4(agent_base)
    agent = _make_file_router4(agent)
    _file_router4_ready = True


# ======== v28h: soft-mode evidence check ========
# Experimental soft-mode evidence check (18 Oct 2026). Needs the v28g agent loaded. Adds a warning line and skips saving when the text an answer was read from has no 'first' or has an extra qualifier. Only the final test loop at the bottom is not part of the patch. NOT in any agent file.
import re, time
QSYN = {"woman":"F","women":"F","female":"F","womens":"F","girl":"F","girls":"F","man":"M","men":"M","male":"M","mens":"M",
        "gold":"gold","silver":"silver","bronze":"bronze","individual":"individual","team":"team","athletics":"athletics",
        "track":"athletics","singles":"singles","doubles":"doubles","mixed":"mixed","junior":"junior"}
def ev_words(s): return re.findall(r"[a-z0-9']+", s.lower().replace("\u2019", "'"))
def ev_quals(ws):
    out = set()
    for w in ws:
        w = w.replace("'s", "").replace("'", "")
        if w in QSYN: out.add(QSYN[w])
    return out
def evidence_supports(question, evidence):
    qw, ew = ev_words(question), ev_words(evidence)
    if "first" in qw:
        if not any(x in ew for x in ("first", "inaugural", "earliest", "1st")):
            return False, "evidence has no 'first'"
        i = ew.index("first") if "first" in ew else 0
        extra = ev_quals(ew[i+1:i+15]) - ev_quals(qw)
        if extra: return False, "extra qualifier in evidence: " + ", ".join(sorted(extra))
    return True, "ok"
def sentences_of(chunk): return [s.strip() for s in re.split(r"(?<=[.!?])\s+|\n", chunk) if len(s.strip()) > 8]
_SKIP = {"The","Indian","India","Olympic","Olympics","Who","Was","First","Test","Academy","Award","Prize","Nobel","Physics","World","Cup","Games","Summer","Ramon","Magsaysay"}
def evidence_for(chunk, answer):
    keys = [t for t in re.findall(r"[A-Z][A-Za-z'.-]{2,}", answer) if t not in _SKIP]
    if not keys: keys = re.findall(r"[a-z]{4,}", answer.lower())
    sents = sentences_of(chunk); ev = []
    for i, s in enumerate(sents):
        if any(k.lower() in s.lower() for k in keys):
            ev.append(s)
            if i + 1 < len(sents): ev.append(sents[i + 1])
    return ev, keys
def shadow_decide(question, chunk, answer):
    ev, keys = evidence_for(chunk, answer)
    if not ev: return "NO-EVIDENCE", "no sentence names " + ", ".join(keys[:3])
    reasons = []
    for s in ev:
        ok, why = evidence_supports(question, s)
        if ok: return "ACCEPT", s[:110]
        reasons.append(why)
    return "REJECT", reasons[0]

captured = []
_orig_ask2 = ask_long
def _cap_ask(*a, **k):
    r = _orig_ask2(*a, **k)
    p = str(a[0]) if a else ""
    if "TEXT:\n" in p and "Does this TEXT directly state" in p:
        try:
            q = p.split("Question: ", 1)[1].split("\n\nTEXT:", 1)[0]
            chunk = p.split("TEXT:\n", 1)[1].split("\n\nDoes this TEXT", 1)[0]
            captured.append({"q": q, "chunk": chunk, "reply": str(r).strip()})
        except Exception: pass
    return r

_agent_all_plain = agent_all
def agent_all(q, *a, **k):
    global ask_long
    before = set(KB.keys()); captured.clear()
    ask_long = _cap_ask
    try:
        out = str(_agent_all_plain(q, *a, **k))
    finally:
        ask_long = _orig_ask2
    yes = [c for c in captured if c["reply"].upper().startswith("YES")]
    pick = None
    for c in reversed(yes):
        ans = c["reply"][4:].strip(" :")
        if any(t.lower() in out.lower() for t in re.findall(r"[A-Za-z]{4,}", ans)): pick = c; break
    if pick is None and yes: pick = yes[-1]
    if pick and "(Source:" in out:
        dec, why = shadow_decide(pick["q"], pick["chunk"], pick["reply"][4:].strip(" :"))
        if dec == "REJECT":
            for key in [x for x in KB.keys() if x not in before]: KB.pop(key, None)
            kb_save(KB)
            out += "\n(Warning: the text this answer was read from does not clearly support it - " + why + ". Please check the source. Not saved.)"
    return out
print("soft-mode check installed")


# ======== v28i: wider code-request rule ========
import re
_CW = getattr(_CODEWORD, "pattern", _CODEWORD)
_START = r"(^|[.;:!?,]\s+)(please\s+|kindly\s+)?(can you\s+|could you\s+|would you\s+)?"
_VERBS = r"(write|create|make|build|generate|fix|debug|run|code|compute|calculate|plot|print|solve|implement|parse|extract|convert|simulate)"
_USING = r"(^|[.;:!?]\s+)using\s+[^,.:;]{1,60},\s+"
_CODE2 = r"\b(pandas|numpy|matplotlib|sqlite3?|dataframe|dijkstra|regex|unittest|flask|django|algorithm)\b|[\w\-]+\.(py|js|ts|sql)\b"
def _more_code(t):
    low = t.lower()
    if re.search(_START + _VERBS + r"\b", low) or re.search(_USING, low):
        if re.search(_CW, low) or re.search(_CODE2, low): return True
    if re.search(_START + r"(compute|calculate)\b", low) and re.search(r"\d", low): return True
    if re.search(r"[:;]\s*print\b", low) and re.search(r"\d", low): return True
    return False
_is_code_prev = is_code_request
def is_code_request(t):
    return bool(_is_code_prev(t)) or _more_code(t)



# ======== v28j: crash fix and safety net ========

import time
def search_titles(query, n):
    for attempt in range(3):
        j = wiki_get({"action": "query", "format": "json", "list": "search",
                      "srsearch": query, "srlimit": n})
        try:
            return [x["title"] for x in j["query"]["search"]]
        except (KeyError, TypeError):
            time.sleep(2 * (attempt + 1))
    return []

_agent_all_safe_prev = agent_all
def agent_all(q, *a, **k):
    try:
        return _agent_all_safe_prev(q, *a, **k)
    except Exception as e:
        try: lang = detect_lang(q)
        except Exception: lang = "English"
        return NOT_FOUND.get(lang, NOT_FOUND["English"]) + "\n(Something went wrong while reading the sources: " + type(e).__name__ + ". Please try again.)"

# ======== v28j: mild note for 'no first' ========

import re as _re
_SOFT_NOFIRST = _re.compile(r"\(Warning: the text this answer was read from does not clearly support it - evidence has no 'first'\. Please check the source\. Not saved\.\)")
_agent_all_before_note = agent_all
def agent_all(q, *a, **k):
    out = _agent_all_before_note(q, *a, **k)
    return _SOFT_NOFIRST.sub("(Note: the source text does not explicitly say 'first'. Not saved.)", out)


# ===== v28l = v28j + K1 (wider code rule) + K2 (run-afterwards net) + K3 corrected (percent hint only when task has % / interest / rate; format-tolerant run comparison) =====

import re
def _more_code2(t):
    low = t.lower().strip()
    cw = re.search(_CW, low) or re.search(_CODE2, low) or re.search(r"\b(recursion|recursive|loop)\b", low)
    concrete = re.search(r"[\d:'\"\[\(]", low)
    if low.startswith(("python:", "python -", "python,")): return True
    if re.match(r"^(i need|i want|i'd like|i would like|help me|can you help me|could you help me)\b", low) and cw: return True
    if re.match(r"^how (do|can|to|would)\b", low) and cw and re.search(r"\b(show|print|output|write|working code)\b", low): return True
    if cw and concrete and re.search(_START + r"(draft|develop|use|sort|count|reverse|remove|convert|check|find|read|get|list)\b", low): return True
    return False
_is_code_prev2 = is_code_request
def is_code_request(t):
    return bool(_is_code_prev2(t)) or _more_code2(t)


import re, tempfile
_PYBLOCK = re.compile(r"```python\s*\n(.*?)```", re.S)
_DENY = re.compile(r"os\.remove|os\.rmdir|rmtree|os\.system|subprocess|socket|input\(|shutil\.move|os\.unlink|__import__|eval\(|exec\(")
_agent_all_before_run = agent_all
def agent_all(q, *a, **k):
    out = str(_agent_all_before_run(q, *a, **k))
    if "Ran the code" in out:
        return out
    m = _PYBLOCK.search(out)
    if not m:
        return out
    code = m.group(1)
    if _DENY.search(code):
        return out + "\n(The code above was not run automatically because it may change files, use the network or ask for input.)"
    rc, res, d = run_python(code, timeout=30)
    if rc == 0 and res.strip():
        return out + "\n\nResult (the code above was run afterwards):\n" + res.strip()
    if res.strip():
        return out + "\n(The code above was run afterwards but failed: " + res.strip().splitlines()[-1][:150] + ")"
    return out


import re as _re3
_PCT3 = _re3.compile(r"%|percent|per cent|interest|\brate\b", _re3.I)
_HINT3 = " (Convert percentages exactly once: 8 percent means 0.08. Check units before printing.)"
def _nums3(s):
    return [float(x) for x in _re3.findall(r"-?\d+(?:\.\d+)?", s)]
def _same3(a, b):
    a, b = a.strip(), b.strip()
    if a == b: return True
    na, nb = _nums3(a), _nums3(b)
    if na and nb:
        if len(na) == len(nb) and all(abs(x - y) <= 0.0051 for x, y in zip(na, nb)): return True
        sa, sb = set(round(x, 2) for x in na), set(round(x, 2) for x in nb)
        return sa <= sb or sb <= sa
    return False
_code_task_prev3 = code_task
def code_task(task, seed=0, tries=3, files=()):
    if files:
        return _code_task_prev3(task, seed, tries, files)
    if _PCT3.search(task) and _HINT3 not in task:
        task = task + _HINT3
    r1 = _code_task_prev3(task, seed, tries, files)
    if not r1.get("ok"):
        return r1
    r2 = _code_task_prev3(task, seed + 1, tries, files)
    if r2.get("ok") and _same3(r1["output"], r2["output"]):
        return r1
    r3 = _code_task_prev3(task, seed + 2, tries, files)
    runs = [r for r in (r1, r2, r3) if r.get("ok")]
    for r in runs:
        group = [x for x in runs if _same3(r["output"], x["output"])]
        if len(group) >= 2:
            pick = group[0]
            others = [x["output"].strip().replace("\n", " | ")[:40] for x in runs if x not in group and _nums3(x["output"])]
            if others:
                pick["output"] = pick["output"].rstrip() + "\n(Note: 2 of 3 runs agreed; the other run gave: " + "; ".join(others) + ".)"
            return pick
    vals = " ; ".join(x["output"].strip().replace("\n", " | ")[:40] for x in runs)
    r1["output"] = r1["output"].rstrip() + "\n(Warning: repeated runs gave different results: " + vals + ". Please check the result yourself.)"
    return r1



# ---- v28o: wider code router + code fallback (19 Oct 2026) ----
def _more_code3(t):
    low = t.lower().strip()
    if re.match(r"^(who|what|when|where|which|why)\b", low): return False
    if re.search(r"\bprint\b", low): return True
    concrete = re.search(r"\d|['\"\[\(]", low)
    if concrete and re.search(_START + r"(count|sort|reverse|remove|transpose|convert|calculate|compute)\b", low): return True
    return False

_CUES = re.compile(r"\b(sum|total|add up|average|mean|median|percent|percentage|power|squared?|squares|cubes?|product|"
                   r"difference|multiples?|divisible|divided|factor|factorial|prime|palindrome|reverse|digits|remainder|"
                   r"smallest|largest|maximum|minimum|seconds|minutes|weekdays|day of the week|uppercase|lowercase|"
                   r"vowels|characters|words|letters|interest|convert|raised|round\w*|sort\w*|area|perimeter|volume|"
                   r"circumference|leap|(even|odd) numbers?|title case|upper case|capitali[sz]e\w*|comma.separated|"
                   r"days (are |is )?(left|between|until|after)|for every)\b", re.I)
_OKCAP = {"i","python","fahrenheit","celsius","kelvin","gst","emi","km","january","february","march","april","may",
          "june","july","august","september","october","november","december","monday","tuesday","wednesday",
          "thursday","friday","saturday","sunday"}
_DATA = re.compile(r"\d|['\"\[\(]|\b\w+, \w+,? (and )?\w+\b|\b(one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|twenty|thirty|hundred|thousand)\b", re.I)
def _needs_outside(t):
    s = re.sub(r"'[^']*'|\"[^\"]*\"", " ", t)
    s = re.sub(r"\d{4}-\d{2}-\d{2}", " ", s)
    if re.search(r"\bin (19|20)\d\d\b", s): return True
    start = True
    for w in re.findall(r"[A-Za-z]+|[.?!]", s):
        if w in ".?!": start = True; continue
        lw = w.lower()
        if w[0].isupper() and not start and lw not in _OKCAP and lw.rstrip("s") not in _OKCAP: return True
        start = False
    return False
def _more_code4(t):
    low = t.lower().strip()
    if re.match(r"^(explain|describe|tell me|why)\b", low): return False
    return bool(_DATA.search(t)) and bool(_CUES.search(t)) and not _needs_outside(t)

def _wrap_router(prev, extra):
    def is_code_request(t):
        return bool(prev(t)) or bool(extra(t))
    return is_code_request
is_code_request = _wrap_router(is_code_request, _more_code3)
is_code_request = _wrap_router(is_code_request, _more_code4)

_FORCE = {"on": False}
def _wrap_force(prev):
    def is_code_request(t):
        return _FORCE["on"] or bool(prev(t))
    return is_code_request
is_code_request = _wrap_force(is_code_request)

_UNVER = re.compile(r"No source checked|sources I found disagree|I could not find it|Only one source found", re.I)
def _wrap_fallback(prev):
    def agent_all(q, *a, **k):
        out = str(prev(q, *a, **k))
        if "Ran the code" in out or not _UNVER.search(out): return out
        if re.match(r"^(explain|describe|tell me|why)\b", q.lower().strip()): return out
        if not _DATA.search(q) or _needs_outside(q): return out
        _FORCE["on"] = True
        try:
            out2 = str(prev(q, *a, **k))
        finally:
            _FORCE["on"] = False
        return out2 if "Ran the code" in out2 else out
    return agent_all
agent_all = _wrap_fallback(agent_all)


# ---- v28p: answer every part of the question (20 Oct 2026) ----
_ALLPARTS = "If the task asks for more than one thing, print each one on its own line with a short label."
if _ALLPARTS not in CODE_PROMPT:
    CODE_PROMPT = CODE_PROMPT.replace("Task: ", _ALLPARTS + " Task: ")


# ---- v28q: percent hint only for numbers like "8 percent"; print each repeated step (20 Oct 2026) ----
_PCT3 = re.compile(r"\d\s*(?:%|percent|per cent)|interest|\brate\b", re.I)
_TRACEQ = "If the task repeats a step several times, print each step number with its value before the final answer."
if _TRACEQ not in CODE_PROMPT:
    CODE_PROMPT = CODE_PROMPT.replace("Task: ", _TRACEQ + " Task: ")
